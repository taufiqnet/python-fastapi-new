from datetime import date
from decimal import Decimal
import io
import openpyxl
import pytest
import pytest_asyncio
from httpx import ASGITransport, AsyncClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app import models_registry  # noqa: F401
from app.database import Base, get_async_db, get_db
from app.main import app
from app.core.identity.models import User
from app.core.tenancy.models import BusinessProfile
from app.core.deps import get_current_user, get_current_user_optional
from app.modules.finance.import_service import SalesImportService, EXPECTED_COLUMNS
from app.modules.finance.models import SalesInvoice
from app.modules.finance.seed import seed_default_chart_of_accounts_sync

sync_engine = create_engine(
    "sqlite:///:memory:",
    connect_args={"check_same_thread": False},
    poolclass=StaticPool,
)
SyncTestingSessionLocal = sessionmaker(
    autocommit=False, autoflush=False, bind=sync_engine
)


@pytest.fixture
def sync_db():
    Base.metadata.create_all(bind=sync_engine)
    db = SyncTestingSessionLocal()

    biz = BusinessProfile(id=1, name_en="Test Import Biz", vat_number="BD123456789", is_active=True)
    db.add(biz)

    user = User(
        id=1,
        username="admin",
        email="admin@example.com",
        password_hash="fakehash",
        is_superuser=True,
        is_active=True,
        business_id=1,
    )
    db.add(user)
    db.commit()

    seed_default_chart_of_accounts_sync(db, 1)

    try:
        yield db
    finally:
        db.close()
        Base.metadata.drop_all(bind=sync_engine)


class AsyncSessionWrapper:
    def __init__(self, db):
        self.db = db

    async def execute(self, query):
        return self.db.execute(query)

    async def commit(self):
        self.db.commit()

    async def flush(self):
        self.db.flush()

    async def refresh(self, obj):
        self.db.refresh(obj)

    def add(self, obj):
        self.db.add(obj)


@pytest_asyncio.fixture
async def client(sync_db):
    def _override_get_db():
        yield sync_db

    async def _override_get_async_db():
        yield AsyncSessionWrapper(sync_db)

    superuser = sync_db.query(User).filter(User.username == "admin").first()

    def _override_get_current_user():
        return superuser

    async def _override_get_current_user_optional(*args, **kwargs):
        return superuser

    app.dependency_overrides[get_db] = _override_get_db
    app.dependency_overrides[get_async_db] = _override_get_async_db
    app.dependency_overrides[get_current_user] = _override_get_current_user
    app.dependency_overrides[get_current_user_optional] = _override_get_current_user_optional

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        yield ac
    app.dependency_overrides.clear()


def test_excel_template_generation():
    template_bytes = SalesImportService.generate_excel_template()
    assert len(template_bytes) > 0

    wb = openpyxl.load_workbook(filename=io.BytesIO(template_bytes))
    ws = wb.active
    headers = [str(c) for c in list(ws.iter_rows(values_only=True))[0]]
    assert headers == EXPECTED_COLUMNS


@pytest.mark.asyncio
async def test_excel_import_validation_and_confirmation(sync_db):
    async_db = AsyncSessionWrapper(sync_db)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(EXPECTED_COLUMNS)
    ws.append([
        "INV-202501-101",
        "2025-01-20",
        "Apex Fabrics Ltd",
        "BD112233445",
        "Gazipur, Dhaka",
        "Gazipur, Dhaka",
        "Covered Van-1",
        "Dyeing Chemical Grade A",
        "KG",
        200,
        150.00,
        0.00,
        15.00,
    ])

    file_bytes = io.BytesIO()
    wb.save(file_bytes)
    file_bytes.seek(0)

    class DummyFile:
        def __init__(self, content, filename):
            self._content = content
            self.filename = filename

        async def read(self):
            return self._content

    file_obj = DummyFile(file_bytes.getvalue(), "sales_sheet.xlsx")

    # 1. Process Upload
    batch = await SalesImportService.process_file_upload(async_db, 1, None, file_obj)
    assert batch.status == "validated"
    assert batch.valid_rows == 1
    assert batch.error_count == 0

    # 2. Confirm Batch Import
    invoices = await SalesImportService.confirm_batch_import(async_db, 1, None, batch.id)
    assert len(invoices) == 1
    assert invoices[0].invoice_number.startswith("INV-202501-")
    assert invoices[0].buyer_name == "Apex Fabrics Ltd"
    assert invoices[0].status == "posted"
    assert invoices[0].voucher_id is not None
