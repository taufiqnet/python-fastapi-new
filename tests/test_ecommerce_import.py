from io import BytesIO
import openpyxl
import pytest
import pytest_asyncio
from httpx import ASGITransport, AsyncClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.database import Base, get_db
from app.main import app

sync_engine = create_engine(
    "sqlite:///:memory:",
    connect_args={"check_same_thread": False},
    poolclass=StaticPool,
)
SyncTestingSessionLocal = sessionmaker(
    autocommit=False, autoflush=False, bind=sync_engine
)


@pytest.fixture
def sync_db():
    Base.metadata.create_all(bind=sync_engine)
    db = SyncTestingSessionLocal()
    try:
        yield db
    finally:
        db.close()
        Base.metadata.drop_all(bind=sync_engine)


@pytest_asyncio.fixture
async def client(sync_db):
    def _override_get_db():
        yield sync_db

    app.dependency_overrides[get_db] = _override_get_db
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        yield ac
    app.dependency_overrides.clear()


@pytest.mark.asyncio
async def test_category_excel_import(client: AsyncClient):
    tpl_res = await client.get("/categories/template-excel?business_id=1")
    assert tpl_res.status_code == 200

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Category Name", "Slug", "Parent Category Name", "Description", "Icon"])
    ws.append(["Home Appliances", "home-appliances", "", "Home appliances category", "fas fa-tv"])
    ws.append(["Refrigerators", "refrigerators", "Home Appliances", "Cooling refrigerators", "fas fa-snowflake"])

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    files = {
        "file": (
            "categories.xlsx",
            output.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }
    import_res = await client.post("/categories/import-excel?business_id=1", files=files)
    assert import_res.status_code == 200
    res_json = import_res.json()
    assert res_json["imported_count"] == 2
    assert len(res_json["errors"]) == 0

    # Duplicate import test
    output.seek(0)
    import_res2 = await client.post("/categories/import-excel?business_id=1", files=files)
    assert import_res2.status_code == 200
    res_json2 = import_res2.json()
    assert res_json2["imported_count"] == 0
    assert len(res_json2["errors"]) == 2


@pytest.mark.asyncio
async def test_brand_excel_import(client: AsyncClient):
    tpl_res = await client.get("/brands/template-excel?business_id=1")
    assert tpl_res.status_code == 200

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Brand Name", "Slug", "Logo URL", "Description"])
    ws.append(["Sony", "sony", "http://example.com/sony.png", "Sony Corporation brand"])

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    files = {
        "file": (
            "brands.xlsx",
            output.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }
    import_res = await client.post("/brands/import-excel?business_id=1", files=files)
    assert import_res.status_code == 200
    res_json = import_res.json()
    assert res_json["imported_count"] == 1

    # Duplicate import test
    output.seek(0)
    import_res2 = await client.post("/brands/import-excel?business_id=1", files=files)
    assert import_res2.status_code == 200
    res_json2 = import_res2.json()
    assert res_json2["imported_count"] == 0
    assert len(res_json2["errors"]) == 1


@pytest.mark.asyncio
async def test_model_excel_import(client: AsyncClient):
    # First create brand
    wb_brand = openpyxl.Workbook()
    ws_brand = wb_brand.active
    ws_brand.append(["Brand Name", "Slug", "Logo URL", "Description"])
    ws_brand.append(["LG", "lg", "", "LG Electronics"])
    out_brand = BytesIO()
    wb_brand.save(out_brand)
    out_brand.seek(0)

    await client.post(
        "/brands/import-excel?business_id=1",
        files={"file": ("brands.xlsx", out_brand.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    tpl_res = await client.get("/brands/models/template-excel?business_id=1")
    assert tpl_res.status_code == 200

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Model Name", "Brand Name", "Slug", "Description"])
    ws.append(["OLED C3", "LG", "oled-c3", "LG OLED TV 2023 Series"])

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    files = {
        "file": (
            "models.xlsx",
            output.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }
    import_res = await client.post("/brands/models/import-excel?business_id=1", files=files)
    assert import_res.status_code == 200
    res_json = import_res.json()
    assert res_json["imported_count"] == 1

    # Duplicate import test
    output.seek(0)
    import_res2 = await client.post("/brands/models/import-excel?business_id=1", files=files)
    assert import_res2.status_code == 200
    res_json2 = import_res2.json()
    assert res_json2["imported_count"] == 0
    assert len(res_json2["errors"]) == 1


@pytest.mark.asyncio
async def test_product_excel_import(client: AsyncClient):
    # Setup Category and Brand
    wb_cat = openpyxl.Workbook()
    ws_cat = wb_cat.active
    ws_cat.append(["Category Name", "Slug", "Parent Category Name", "Description", "Icon"])
    ws_cat.append(["Laptops", "laptops", "", "Computers", "fas fa-laptop"])
    out_cat = BytesIO()
    wb_cat.save(out_cat)
    out_cat.seek(0)
    await client.post(
        "/categories/import-excel?business_id=1",
        files={"file": ("cat.xlsx", out_cat.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    wb_brand = openpyxl.Workbook()
    ws_brand = wb_brand.active
    ws_brand.append(["Brand Name", "Slug", "Logo URL", "Description"])
    ws_brand.append(["Dell", "dell", "", "Dell Technologies"])
    out_brand = BytesIO()
    wb_brand.save(out_brand)
    out_brand.seek(0)
    await client.post(
        "/brands/import-excel?business_id=1",
        files={"file": ("brand.xlsx", out_brand.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    tpl_res = await client.get("/products/template-excel?business_id=1")
    assert tpl_res.status_code == 200

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([
        "Product Title", "Slug", "Category Name", "Brand Name", "Model Name",
        "Status (draft/published/archived)", "Condition (new/used/refurbished/open_box)",
        "Product Type (physical/digital/service)", "Price", "SKU", "Stock Quantity", "Description"
    ])
    ws.append([
        "Dell XPS 15 Laptop", "dell-xps-15-laptop", "Laptops", "Dell", "",
        "published", "new", "physical", 1499.99, "DELL-XPS15-01", 20, "High performance laptop"
    ])

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    files = {
        "file": (
            "products.xlsx",
            output.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }
    import_res = await client.post("/products/import-excel?business_id=1", files=files)
    assert import_res.status_code == 200
    res_json = import_res.json()
    assert res_json["imported_count"] == 1

    # Duplicate import test
    output.seek(0)
    import_res2 = await client.post("/products/import-excel?business_id=1", files=files)
    assert import_res2.status_code == 200
    res_json2 = import_res2.json()
    assert res_json2["imported_count"] == 0
    assert len(res_json2["errors"]) == 1
