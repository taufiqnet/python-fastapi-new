import pytest
import pytest_asyncio
from httpx import ASGITransport, AsyncClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.database import Base, get_db
from app.main import app

sync_engine = create_engine(
    "sqlite:///:memory:",
    connect_args={"check_same_thread": False},
    poolclass=StaticPool,
)
SyncTestingSessionLocal = sessionmaker(
    autocommit=False, autoflush=False, bind=sync_engine
)


@pytest.fixture
def sync_db():
    Base.metadata.create_all(bind=sync_engine)
    db = SyncTestingSessionLocal()
    try:
        yield db
    finally:
        db.close()
        Base.metadata.drop_all(bind=sync_engine)


from app.core.deps import get_current_user_optional, get_current_user
from app.core.identity.models import User

@pytest_asyncio.fixture
async def client(sync_db):
    def _override_get_db():
        yield sync_db

    from app.modules.hr_payroll.employees.router import employee_service
    orig_bg_func = employee_service.process_import_job_background

    def _sync_bg_wrapper(*args, **kwargs):
        kwargs["session_factory"] = lambda: sync_db
        orig_bg_func(*args, **kwargs)

    employee_service.process_import_job_background = _sync_bg_wrapper

    dummy_user = User(
        id=1,
        username="admin",
        email="admin@example.com",
        is_active=True,
        is_superuser=True,
        business_id=1,
    )
    def _override_get_current_user(request=None):
        return dummy_user

    app.dependency_overrides[get_db] = _override_get_db
    app.dependency_overrides[get_current_user_optional] = _override_get_current_user
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        yield ac
    app.dependency_overrides.clear()


@pytest.mark.asyncio
async def test_employee_crud_and_validation(client: AsyncClient):
    # 1. Create Department & Job Title
    dept_res = await client.post(
        "/departments",
        json={
            "name": "Engineering",
            "slug": "engineering",
            "description": "Eng Dept",
            "is_active": True,
            "multiple_heads_allowed": False,
            "business_id": 1,
        },
    )
    assert dept_res.status_code == 201
    dept_id = dept_res.json()["id"]

    jt_res = await client.post(
        "/job-titles",
        json={
            "name": "Software Engineer",
            "short_name": "SE",
            "description": "Dev",
            "is_active": True,
            "department_id": dept_id,
            "business_id": 1,
        },
    )
    assert jt_res.status_code == 201
    jt_id = jt_res.json()["id"]

    # 2. Create Employee 1 (Department Head)
    emp1_payload = {
        "first_name": "John",
        "last_name": "Doe",
        "employee_id": "EMP-001",
        "work_email": "john.doe@example.com",
        "phone": "+1234567890",
        "department_id": dept_id,
        "job_title_id": jt_id,
        "is_department_head": True,
        "is_active": True,
        "business_id": 1,
    }
    res = await client.post("/employees", json=emp1_payload)
    assert res.status_code == 201, res.text
    emp1_data = res.json()
    assert emp1_data["full_name"] == "John Doe"
    assert emp1_data["is_department_head"] is True
    emp1_id = emp1_data["id"]

    # 3. Duplicate Employee ID (should fail)
    dup_payload = emp1_payload.copy()
    dup_payload["work_email"] = "different@example.com"
    dup_payload["phone"] = "+1999999999"
    res_dup = await client.post("/employees", json=dup_payload)
    assert res_dup.status_code == 400

    # 4. Attempt second Department Head in single-head department (should fail)
    emp2_payload = {
        "first_name": "Jane",
        "last_name": "Smith",
        "employee_id": "EMP-002",
        "work_email": "jane.smith@example.com",
        "phone": "+1987654321",
        "department_id": dept_id,
        "job_title_id": jt_id,
        "direct_manager_id": emp1_id,
        "is_department_head": True,
        "is_active": True,
        "business_id": 1,
    }
    res_head_fail = await client.post("/employees", json=emp2_payload)
    assert res_head_fail.status_code == 400

    # 5. Create Employee 2 as regular employee under Employee 1
    emp2_payload["is_department_head"] = False
    res2 = await client.post("/employees", json=emp2_payload)
    assert res2.status_code == 201
    emp2_id = res2.json()["id"]

    # 6. Get List & Detail
    res_get = await client.get(f"/employees/{emp1_id}")
    assert res_get.status_code == 200
    assert res_get.json()["id"] == emp1_id

    res_list = await client.get(f"/employees?business_id=1&department_id={dept_id}")
    assert res_list.status_code == 200
    assert len(res_list.json()) == 2

    # 7. Update Employee 2
    res_upd = await client.put(
        f"/employees/{emp2_id}", json={"middle_name": "Ann"}
    )
    assert res_upd.status_code == 200
    assert res_upd.json()["full_name"] == "Jane Ann Smith"

    # 8. Delete Employees
    res_del2 = await client.delete(f"/employees/{emp2_id}")
    assert res_del2.status_code == 204

    res_del1 = await client.delete(f"/employees/{emp1_id}")
    assert res_del1.status_code == 204

    res_verify = await client.get(f"/employees/{emp1_id}")
    assert res_verify.status_code == 404


@pytest.mark.asyncio
async def test_employee_excel_template_and_import(client: AsyncClient):
    # Template download
    tpl_res = await client.get("/employees/template-excel?business_id=1")
    assert tpl_res.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in tpl_res.headers["content-type"]

    # Import test file
    import openpyxl
    from io import BytesIO

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([
        "Employee ID", "First Name", "Middle Name", "Last Name", "Work Email",
        "Phone", "Department", "Job Title", "Employment Type", "Work Arrangement",
        "Start Date", "Gender"
    ])
    ws.append([
        "IMP100", "Alice", "", "Wonderland", "alice@example.com",
        "+111222333", "", "", "full_time", "remote",
        "2025-01-15", "female"
    ])
    output = BytesIO()
    wb.save(output)
    output.seek(0)

    files = {
        "file": (
            "employees.xlsx",
            output.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }
    import_res = await client.post("/employees/import-excel?business_id=1", files=files)
    assert import_res.status_code == 202
    res_json = import_res.json()
    assert "job_id" in res_json
    assert res_json["status"] == "processing"
    job_id = res_json["job_id"]

    # Poll status endpoint
    status_res = await client.get(f"/employees/import/{job_id}/status")
    assert status_res.status_code == 200
    status_data = status_res.json()
    assert status_data["job_id"] == job_id
    assert status_data["total_rows"] == 1

    # Test error CSV endpoint
    err_csv_res = await client.get(f"/employees/import/{job_id}/errors-csv")
    assert err_csv_res.status_code == 200
    assert "text/csv" in err_csv_res.headers["content-type"]

    # Test cancel endpoint
    cancel_res = await client.post(f"/employees/import/{job_id}/cancel")
    assert cancel_res.status_code == 200
    assert cancel_res.json()["cancelled"] is True


@pytest.mark.asyncio
async def test_employee_import_tenant_scoping_superuser_and_non_superuser(client: AsyncClient, sync_db):
    import openpyxl
    from io import BytesIO
    from app.core.tenancy.models import BusinessProfile

    # Create business 2
    biz2 = BusinessProfile(
        id=2,
        legal_name="Dreamlight",
        name_en="Dreamlight",
        cr_number="9876543210",
        vat_number="300000000000009",
    )
    sync_db.add(biz2)
    sync_db.commit()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([
        "Employee ID", "First Name", "Middle Name", "Last Name", "Work Email",
        "Phone", "Department", "Job Title", "Employment Type", "Work Arrangement",
        "Start Date", "Gender"
    ])
    ws.append([
        "SUPER01", "Super", "", "UserEmp", "super.emp@dreamlight.com",
        "+1112223334", "", "", "full_time", "remote",
        "2025-01-15", "female"
    ])
    output = BytesIO()
    wb.save(output)
    excel_bytes = output.getvalue()

    # 1. Test Superuser import specifying business_id=2
    files1 = {"file": ("super_import.xlsx", excel_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    res1 = await client.post("/employees/import-excel?business_id=2", files=files1)
    assert res1.status_code == 202
    job_id1 = res1.json()["job_id"]

    # Check imported employees under business_id=2
    emp_res1 = await client.get("/employees?business_id=2")
    assert emp_res1.status_code == 200
    emps1 = emp_res1.json()
    assert len(emps1) == 1
    assert emps1[0]["employee_id"] == "SUPER01"
    assert emps1[0]["business_id"] == 2

    # 2. Test Non-superuser (assigned to business_id=1) submitting business_id=2 in request
    non_super_user = User(
        id=2,
        username="regular_hr",
        email="hr@mycompany.com",
        is_active=True,
        is_superuser=False,
        business_id=1,
    )
    non_super_user.get_all_permission_codes = lambda: {"hrm:employees:create", "hrm:employees:view"}
    non_super_user.has_permission = lambda code: code in {"hrm:employees:create", "hrm:employees:view"}
    non_super_user.business_profile = None

    app.dependency_overrides[get_current_user_optional] = lambda request=None: non_super_user
    app.dependency_overrides[get_current_user] = lambda request=None: non_super_user

    wb2 = openpyxl.Workbook()
    ws2 = wb2.active
    ws2.append([
        "Employee ID", "First Name", "Middle Name", "Last Name", "Work Email",
        "Phone", "Department", "Job Title", "Employment Type", "Work Arrangement",
        "Start Date", "Gender"
    ])
    ws2.append([
        "NONSUP01", "NonSuper", "", "Emp", "nonsuper.emp@mycompany.com",
        "+1112223335", "", "", "full_time", "onsite",
        "2025-01-15", "male"
    ])
    out2 = BytesIO()
    wb2.save(out2)
    excel_bytes2 = out2.getvalue()

    files2 = {"file": ("nonsuper_import.xlsx", excel_bytes2, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    res2 = await client.post("/employees/import-excel?business_id=2", files=files2)
    assert res2.status_code == 202
    job_id2 = res2.json()["job_id"]

    # Verify job ownership / status
    job_status_res = await client.get(f"/employees/import/{job_id2}/status")
    assert job_status_res.status_code == 200
    assert job_status_res.json()["success_count"] == 1

    # Assert imported employee landed in business_id=1 (non-superuser's own business), NOT business_id=2
    emp_res2_biz1 = await client.get("/employees?business_id=1")
    assert emp_res2_biz1.status_code == 200
    biz1_emps = [e for e in emp_res2_biz1.json() if e["employee_id"] == "NONSUP01"]
    assert len(biz1_emps) == 1
    assert biz1_emps[0]["business_id"] == 1

    emp_res2_biz2 = await client.get("/employees?business_id=2")
    assert emp_res2_biz2.status_code == 200
    biz2_nonsup_emps = [e for e in emp_res2_biz2.json() if e["employee_id"] == "NONSUP01"]
    assert len(biz2_nonsup_emps) == 0
