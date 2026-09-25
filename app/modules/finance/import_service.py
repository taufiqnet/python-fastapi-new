import io
import uuid
from datetime import date, datetime
from decimal import Decimal
from typing import Any

import openpyxl
from fastapi import HTTPException, UploadFile, status
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.modules.finance.models import Customer, SalesImportBatch, SalesInvoice
from app.modules.finance.schemas import SalesInvoiceCreate, SalesInvoiceLineCreate
from app.modules.finance.services import InvoiceService, Mushak63Service


EXPECTED_COLUMNS = [
    "Invoice Number",
    "Invoice Date",
    "Customer Name",
    "Customer BIN",
    "Customer Address",
    "Delivery Destination",
    "Vehicle Info",
    "Item Description",
    "UOM",
    "Quantity",
    "Unit Price",
    "SD Rate (%)",
    "VAT Rate (%)",
]


class SalesImportService:
    @staticmethod
    def generate_excel_template() -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sales Import Template"

        # Headers
        ws.append(EXPECTED_COLUMNS)

        # Sample rows
        sample_rows = [
            [
                "INV-202501-001",
                "2025-01-15",
                "Dhaka Enterprise Ltd",
                "BD123456789",
                "Gulshan 1, Dhaka",
                "Dhanmondi, Dhaka",
                "Truck DHAKA-METRO-11",
                "Cotton Fabric Roll",
                "Meter",
                100,
                250.00,
                0.00,
                15.00,
            ],
            [
                "INV-202501-001",
                "2025-01-15",
                "Dhaka Enterprise Ltd",
                "BD123456789",
                "Gulshan 1, Dhaka",
                "Dhanmondi, Dhaka",
                "Truck DHAKA-METRO-11",
                "Polyester Yarn",
                "KG",
                50,
                180.00,
                5.00,
                15.00,
            ],
            [
                "INV-202501-002",
                "2025-01-16",
                "Chittagong Trading",
                "",
                "Agrabad, Chittagong",
                "Agrabad, Chittagong",
                "",
                "Sewing Thread",
                "Spool",
                200,
                45.00,
                0.00,
                15.00,
            ],
        ]

        for row in sample_rows:
            ws.append(row)

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    @staticmethod
    async def process_file_upload(
        db: AsyncSession, business_id: int, user_id: int | None, file: UploadFile
    ) -> SalesImportBatch:
        content = await file.read()
        if not content:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST, detail="Uploaded file is empty."
            )

        filename = file.filename or "sales_import.xlsx"
        is_csv = filename.endswith(".csv")

        parsed_rows: list[dict[str, Any]] = []

        if is_csv:
            import csv

            csv_text = content.decode("utf-8-sig")
            reader = csv.DictReader(io.StringIO(csv_text))
            for r in reader:
                parsed_rows.append(r)
        else:
            wb = openpyxl.load_workbook(filename=io.BytesIO(content), data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if not rows or len(rows) < 2:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="File must contain a header row and at least one data row.",
                )

            headers = [str(c).strip() if c else "" for c in rows[0]]
            for row in rows[1:]:
                if not any(row):
                    continue
                row_dict = {}
                for idx, val in enumerate(row):
                    if idx < len(headers):
                        row_dict[headers[idx]] = str(val).strip() if val is not None else ""
                parsed_rows.append(row_dict)

        # Validate rows
        row_errors = []
        valid_count = 0

        # Existing invoices check
        existing_res = await db.execute(
            select(SalesInvoice.invoice_number).where(SalesInvoice.business_id == business_id)
        )
        existing_inv_numbers = set(existing_res.scalars().all())

        grouped_invoice_buyers: dict[str, str] = {}

        for idx, row in enumerate(parsed_rows, start=2):
            inv_num = row.get("Invoice Number") or row.get("invoice_number") or ""
            inv_date_str = row.get("Invoice Date") or row.get("invoice_date") or ""
            buyer_name = row.get("Customer Name") or row.get("customer_name") or ""
            item_desc = row.get("Item Description") or row.get("item_description") or ""
            qty_str = row.get("Quantity") or row.get("quantity") or ""
            price_str = row.get("Unit Price") or row.get("unit_price") or ""
            sd_str = row.get("SD Rate (%)") or row.get("sd_rate") or "0"
            vat_str = row.get("VAT Rate (%)") or row.get("vat_rate") or "15"

            errors = []
            if not inv_num:
                errors.append("Invoice Number is required.")
            elif inv_num in existing_inv_numbers:
                errors.append(f"Invoice Number '{inv_num}' already exists in database.")

            if not inv_date_str:
                errors.append("Invoice Date is required.")
            else:
                try:
                    # Parse date format YYYY-MM-DD
                    datetime.strptime(inv_date_str[:10], "%Y-%m-%d")
                except Exception:
                    errors.append(f"Invalid date format '{inv_date_str}'. Expected YYYY-MM-DD.")

            if not buyer_name:
                errors.append("Customer Name is required.")

            if inv_num and buyer_name:
                if inv_num in grouped_invoice_buyers and grouped_invoice_buyers[inv_num] != buyer_name:
                    errors.append(f"Inconsistent Customer Name for Invoice Number '{inv_num}'.")
                else:
                    grouped_invoice_buyers[inv_num] = buyer_name

            if not item_desc:
                errors.append("Item Description is required.")

            try:
                qty = float(qty_str)
                if qty <= 0:
                    errors.append("Quantity must be greater than 0.")
            except Exception:
                errors.append(f"Invalid Quantity '{qty_str}'.")

            try:
                price = float(price_str)
                if price < 0:
                    errors.append("Unit Price cannot be negative.")
            except Exception:
                errors.append(f"Invalid Unit Price '{price_str}'.")

            try:
                sd = float(sd_str)
                if sd < 0 or sd > 100:
                    errors.append("SD Rate (%) must be between 0 and 100.")
            except Exception:
                errors.append(f"Invalid SD Rate '{sd_str}'.")

            try:
                vat = float(vat_str)
                if vat < 0 or vat > 100:
                    errors.append("VAT Rate (%) must be between 0 and 100.")
            except Exception:
                errors.append(f"Invalid VAT Rate '{vat_str}'.")

            if errors:
                row_errors.append({
                    "row": idx,
                    "invoice_number": inv_num,
                    "errors": errors,
                })
            else:
                valid_count += 1

        batch_status = "validated" if not row_errors else "failed"

        batch = SalesImportBatch(
            business_id=business_id,
            filename=filename,
            status=batch_status,
            total_rows=len(parsed_rows),
            valid_rows=valid_count,
            error_count=len(row_errors),
            error_summary={"errors": row_errors} if row_errors else None,
            raw_data={"rows": parsed_rows},
            created_by_id=user_id,
        )
        db.add(batch)
        await db.commit()
        await db.refresh(batch)
        return batch

    @staticmethod
    async def confirm_batch_import(
        db: AsyncSession, business_id: int, user_id: int | None, batch_id: uuid.UUID
    ) -> list[SalesInvoice]:
        res = await db.execute(
            select(SalesImportBatch)
            .where(SalesImportBatch.business_id == business_id)
            .where(SalesImportBatch.id == batch_id)
        )
        batch = res.scalar_one_or_none()
        if not batch:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND, detail="Sales import batch not found."
            )

        if batch.status == "confirmed":
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST, detail="Batch is already confirmed."
            )

        if batch.status == "failed" or batch.error_count > 0:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Cannot confirm batch with {batch.error_count} validation errors. Fix errors and re-upload.",
            )

        rows = (batch.raw_data or {}).get("rows", [])
        if not rows:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST, detail="No rows found in batch data."
            )

        # Group rows by invoice number
        grouped_invoices: dict[str, list[dict[str, Any]]] = {}
        for r in rows:
            inv_num = r.get("Invoice Number") or r.get("invoice_number")
            grouped_invoices.setdefault(inv_num, []).append(r)

        created_invoices: list[SalesInvoice] = []

        for inv_num, rlist in grouped_invoices.items():
            first = rlist[0]
            buyer_name = first.get("Customer Name") or first.get("customer_name")
            buyer_bin = first.get("Customer BIN") or first.get("customer_bin") or None
            buyer_address = first.get("Customer Address") or first.get("customer_address") or None
            del_dest = first.get("Delivery Destination") or first.get("delivery_destination") or None
            veh_info = first.get("Vehicle Info") or first.get("vehicle_info") or None
            inv_date_str = first.get("Invoice Date") or first.get("invoice_date")
            inv_date = datetime.strptime(inv_date_str[:10], "%Y-%m-%d").date()

            # Auto-match or create Customer
            cust_query = select(Customer).where(Customer.business_id == business_id)
            if buyer_bin:
                cust_query = cust_query.where(Customer.bin == buyer_bin)
            else:
                cust_query = cust_query.where(Customer.name == buyer_name)

            cust_res = await db.execute(cust_query)
            cust = cust_res.scalar_one_or_none()

            if not cust:
                cust = Customer(
                    business_id=business_id,
                    name=buyer_name,
                    bin=buyer_bin,
                    address=buyer_address,
                )
                db.add(cust)
                await db.flush()

            lines_create = []
            for r in rlist:
                item_desc = r.get("Item Description") or r.get("item_description")
                uom = r.get("UOM") or r.get("uom") or "Pcs"
                qty = Decimal(str(r.get("Quantity") or r.get("quantity")))
                uprice = Decimal(str(r.get("Unit Price") or r.get("unit_price")))
                sd_rate = Decimal(str(r.get("SD Rate (%)") or r.get("sd_rate") or "0"))
                vat_rate = Decimal(str(r.get("VAT Rate (%)") or r.get("vat_rate") or "15"))

                lines_create.append(
                    SalesInvoiceLineCreate(
                        description=item_desc,
                        uom=uom,
                        quantity=qty,
                        unit_price=uprice,
                        sd_rate=sd_rate,
                        vat_rate=vat_rate,
                    )
                )

            inv_data = SalesInvoiceCreate(
                business_id=business_id,
                issue_date=inv_date,
                customer_id=cust.id,
                buyer_name=buyer_name,
                buyer_bin=buyer_bin,
                buyer_address=buyer_address,
                delivery_destination=del_dest,
                vehicle_nature_number=veh_info,
                lines=lines_create,
            )

            invoice = await InvoiceService.create_invoice(db, business_id, user_id, inv_data)
            # Auto-post invoice journal entry
            posted_inv = await InvoiceService.post_invoice(db, business_id, user_id, invoice.id)

            # Issue Mushak 6.3 Tax Invoice so imported sales invoices immediately appear in the Mushak list
            try:
                await Mushak63Service.issue_mushak_challan(db, business_id, user_id, posted_inv.id)
            except Exception:
                pass

            created_invoices.append(posted_inv)

        batch.status = "confirmed"
        await db.commit()
        return created_invoices
