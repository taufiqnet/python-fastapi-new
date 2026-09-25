import uuid
from decimal import Decimal

from fastapi import HTTPException, status
from sqlalchemy.orm import Session

from app.modules.ecommerce.inventory.models import (
    StockMovementReason,
)
from app.modules.ecommerce.inventory.schemas import StockAdjustmentRequest
from app.modules.ecommerce.inventory.service import InventoryService
from app.modules.ecommerce.orders.models import (
    OrderFulfillmentStatus,
    OrderPaymentStatus,
)
from app.modules.ecommerce.orders.repository import OrderRepository
from app.modules.ecommerce.orders.schemas import (
    OrderCreate,
    OrderDetail,
    OrderStatusUpdate,
    OrderSummary,
    OrderUpdate,
)
from app.modules.ecommerce.products.models import Product, ProductVariant


class OrderService:
    def __init__(
        self,
        repository: OrderRepository | None = None,
        inventory_service: InventoryService | None = None,
    ):
        self.repository = repository or OrderRepository()
        self.inventory_service = inventory_service or InventoryService()

    def _process_inventory_deduction_for_order(self, db: Session, order) -> None:
        """Deducts inventory stock for each order line item when status is completed or delivered."""
        # Find default warehouse for the business or fallback to any warehouse
        warehouses = self.inventory_service.get_warehouses(
            db, business_id=order.business_id, is_active=True
        )
        if not warehouses:
            # Fallback to default or any active warehouse
            warehouses = self.inventory_service.get_warehouses(db, is_active=True)
        
        if not warehouses:
            # If no warehouse exists, create a default warehouse for the business profile
            from app.modules.ecommerce.inventory.schemas import WarehouseCreate
            wh = self.inventory_service.create_warehouse(
                db,
                WarehouseCreate(
                    name="Main Warehouse",
                    code=f"MAIN-WH-{order.business_id}",
                    business_id=order.business_id,
                    is_active=True,
                    is_default=True,
                ),
            )
            warehouse_id = wh.id
        else:
            default_wh = next((w for w in warehouses if w.is_default), warehouses[0])
            warehouse_id = default_wh.id

        for item in order.items:
            # Check or create inventory item for the variant
            inv_items = self.inventory_service.repository.get_inventory_items(
                db, variant_id=item.variant_id, warehouse_id=warehouse_id
            )
            if inv_items:
                inv_item = inv_items[0]
            else:
                from app.modules.ecommerce.inventory.schemas import InventoryItemCreate
                inv_item = self.inventory_service.repository.create_inventory_item(
                    db,
                    InventoryItemCreate(
                        variant_id=item.variant_id,
                        warehouse_id=warehouse_id,
                        quantity_on_hand=item.quantity,  # set initial stock before deduction or 0
                    ),
                )

            # Record stock deduction movement
            adj_data = StockAdjustmentRequest(
                inventory_item_id=inv_item.id,
                delta=-item.quantity,
                reason=StockMovementReason.ORDER,
                reference_id=str(order.id),
                notes=f"Stock deducted for completed order {order.order_number}",
            )
            self.inventory_service.adjust_stock(db, adj_data)

    def create_order(self, db: Session, data: OrderCreate) -> OrderDetail:
        calculated_total = Decimal("0.00")
        items_with_prices = []

        for item in data.items:
            variant = (
                db.query(ProductVariant)
                .filter(ProductVariant.id == item.variant_id)
                .first()
            )
            if not variant:
                raise HTTPException(
                    status_code=status.HTTP_404_NOT_FOUND,
                    detail=f"Variant with ID {item.variant_id} not found",
                )

            product = db.query(Product).filter(Product.id == variant.product_id).first()
            seller_id = product.seller_id if product else None

            unit_price = variant.price
            subtotal = unit_price * item.quantity
            calculated_total += subtotal

            items_with_prices.append(
                (
                    variant.id,
                    seller_id,
                    product.title if product else "Product",
                    variant.sku,
                    variant.attributes,
                    item.quantity,
                    unit_price,
                    subtotal,
                )
            )

        order = self.repository.create_order(
            db,
            data=data,
            calculated_total=calculated_total,
            items_with_prices=items_with_prices,
        )
        return OrderDetail.model_validate(order)

    def get_order(
        self, db: Session, order_id: uuid.UUID, business_id: int = 1
    ) -> OrderDetail:
        order = self.repository.get_order_by_id(db, order_id, business_id=business_id)
        if not order:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND, detail="Order not found"
            )
        return OrderDetail.model_validate(order)

    def get_orders(
        self,
        db: Session,
        business_id: int = 1,
        user_id: uuid.UUID | None = None,
        payment_status: OrderPaymentStatus | None = None,
        fulfillment_status: OrderFulfillmentStatus | None = None,
        skip: int = 0,
        limit: int = 100,
    ) -> list[OrderSummary]:
        orders = self.repository.get_orders(
            db,
            business_id=business_id,
            user_id=user_id,
            payment_status=payment_status,
            fulfillment_status=fulfillment_status,
            skip=skip,
            limit=limit,
        )
        return [OrderSummary.model_validate(o) for o in orders]

    def update_order_status(
        self,
        db: Session,
        order_id: uuid.UUID,
        update_data: OrderStatusUpdate,
        business_id: int = 1,
    ) -> OrderDetail:
        order = self.repository.get_order_by_id(db, order_id, business_id=business_id)
        if not order:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND, detail="Order not found"
            )
        
        prev_f_status = getattr(order.fulfillment_status, "value", order.fulfillment_status)
        updated = self.repository.update_order_status(db, order, update_data)
        new_f_status = getattr(updated.fulfillment_status, "value", updated.fulfillment_status)

        # Trigger inventory deduction if status transitions to completed or delivered
        if new_f_status in ("delivered", "completed") and prev_f_status not in ("delivered", "completed"):
            self._process_inventory_deduction_for_order(db, updated)

        # Trigger automatic Mushak 6.3 generation if Time of Supply rules are met
        try:
            from app.modules.finance.services import generateMushak63, Mushak63Service
            eligible, _ = Mushak63Service.is_order_mushak_eligible(updated, db_sync=db)
            if eligible:
                generateMushak63(updated.id, db=db, business_id=updated.business_id)
        except Exception as e:
            # Non-blocking log or pass if finance module isn't configured or errors
            pass

        return OrderDetail.model_validate(updated)

    def update_order(
        self,
        db: Session,
        order_id: uuid.UUID,
        data: OrderUpdate,
        business_id: int = 1,
    ) -> OrderDetail:
        order = self.repository.get_order_by_id(db, order_id, business_id=business_id)
        if not order:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND, detail="Order not found"
            )

        calculated_total = None
        items_with_prices = None

        if data.items is not None:
            calculated_total = Decimal("0.00")
            items_with_prices = []

            for item in data.items:
                variant = (
                    db.query(ProductVariant)
                    .filter(ProductVariant.id == item.variant_id)
                    .first()
                )
                if not variant:
                    raise HTTPException(
                        status_code=status.HTTP_404_NOT_FOUND,
                        detail=f"Variant with ID {item.variant_id} not found",
                    )

                product = db.query(Product).filter(Product.id == variant.product_id).first()
                seller_id = product.seller_id if product else None

                unit_price = variant.price
                subtotal = unit_price * item.quantity
                calculated_total += subtotal

                items_with_prices.append(
                    (
                        variant.id,
                        seller_id,
                        product.title if product else "Product",
                        variant.sku,
                        variant.attributes,
                        item.quantity,
                        unit_price,
                        subtotal,
                    )
                )

        prev_f_status = getattr(order.fulfillment_status, "value", order.fulfillment_status)
        updated = self.repository.update_order(
            db,
            order=order,
            data=data,
            calculated_total=calculated_total,
            items_with_prices=items_with_prices,
        )
        new_f_status = getattr(updated.fulfillment_status, "value", updated.fulfillment_status)

        # Trigger inventory deduction if status transitions to completed or delivered
        if new_f_status in ("delivered", "completed") and prev_f_status not in ("delivered", "completed"):
            self._process_inventory_deduction_for_order(db, updated)

        # Trigger automatic Mushak 6.3 generation if Time of Supply rules are met
        try:
            from app.modules.finance.services import generateMushak63, Mushak63Service
            eligible, _ = Mushak63Service.is_order_mushak_eligible(updated, db_sync=db)
            if eligible:
                generateMushak63(updated.id, db=db, business_id=updated.business_id)
        except Exception as e:
            pass

        return OrderDetail.model_validate(updated)

    def delete_order(
        self,
        db: Session,
        order_id: uuid.UUID,
        business_id: int = 1,
    ) -> None:
        order = self.repository.get_order_by_id(db, order_id, business_id=business_id)
        if not order:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND, detail="Order not found"
            )
        self.repository.delete_order(db, order)

    def generate_excel_template(self, db: Session, business_id: int) -> bytes:
        import io

        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Orders"
        ws.append([
            "Recipient Name", "Phone", "Email", "Street", "City", "State", "Zip Code", "Country",
            "Product SKU", "Quantity", "Currency", "Shipping Fee", "Tax Amount", "Discount Amount",
            "Other Charges", "Payment Method", "Payment Channel", "Delivery Method", "Courier Company",
            "Tracking ID", "Customer Note", "Payment Status", "Fulfillment Status"
        ])
        ws.append([
            "Sample Customer", "+8801700000000", "customer@example.com", "123 Main St", "Dhaka", "Dhaka", "1205", "Bangladesh",
            "SKU-12345", 2, "BDT", 60.0, 0.0, 0.0, 0.0, "cash_on_delivery", "Bkash", "Inside Dhaka", "Pathao",
            "TRACK-12345", "Please call before delivery", "unpaid", "pending"
        ])
        excel_file = io.BytesIO()
        wb.save(excel_file)
        return excel_file.getvalue()

    def import_orders_excel(self, db: Session, business_id: int, file_bytes: bytes) -> dict:
        import io

        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(file_bytes))
        ws = wb.active
        imported_count = 0
        errors = []

        rows = list(ws.iter_rows(values_only=True))
        if not rows or len(rows) < 2:
            return {"imported_count": 0, "errors": ["Excel file is empty or missing headers"]}

        # Dynamic header-based column mapping
        header_map = {}
        header_row = rows[0]
        if header_row:
            for c_idx, cell in enumerate(header_row):
                if cell is None:
                    continue
                clean_h = str(cell).strip().lower().replace("_", " ").replace("-", " ")
                if clean_h in ("recipient name", "recipient", "customer name", "customer"):
                    header_map["recipient_name"] = c_idx
                elif clean_h in ("phone", "phone number", "mobile"):
                    header_map["phone"] = c_idx
                elif clean_h in ("email", "email address"):
                    header_map["email"] = c_idx
                elif clean_h in ("street", "street address", "address"):
                    header_map["street"] = c_idx
                elif clean_h == "city":
                    header_map["city"] = c_idx
                elif clean_h in ("state", "province"):
                    header_map["state"] = c_idx
                elif clean_h in ("zip code", "zip", "postal code", "zipcode"):
                    header_map["zip_code"] = c_idx
                elif clean_h == "country":
                    header_map["country"] = c_idx
                elif clean_h in ("product sku", "sku", "variant sku", "item sku"):
                    header_map["sku"] = c_idx
                elif clean_h in ("quantity", "qty", "count"):
                    header_map["quantity"] = c_idx
                elif clean_h == "currency":
                    header_map["currency"] = c_idx
                elif clean_h in ("shipping fee", "shipping amount", "shipping"):
                    header_map["shipping_fee"] = c_idx
                elif clean_h in ("tax amount", "tax"):
                    header_map["tax_amount"] = c_idx
                elif clean_h in ("discount amount", "discount", "campaign discount"):
                    header_map["discount_amount"] = c_idx
                elif clean_h in ("other charges", "other charge", "charges"):
                    header_map["other_charges"] = c_idx
                elif clean_h == "payment method":
                    header_map["payment_method"] = c_idx
                elif clean_h == "payment channel":
                    header_map["payment_channel"] = c_idx
                elif clean_h == "delivery method":
                    header_map["delivery_method"] = c_idx
                elif clean_h in ("courier company", "courier"):
                    header_map["courier_company"] = c_idx
                elif clean_h in ("tracking id", "tracking number"):
                    header_map["tracking_id"] = c_idx
                elif clean_h in ("customer note", "note", "notes", "instructions"):
                    header_map["customer_note"] = c_idx
                elif clean_h == "payment status":
                    header_map["payment_status"] = c_idx
                elif clean_h == "fulfillment status":
                    header_map["fulfillment_status"] = c_idx

        def get_val(r_cells, field_key, pos_fallback=None):
            if field_key in header_map and header_map[field_key] < len(r_cells):
                val = r_cells[header_map[field_key]]
                if val is not None:
                    return val
            if pos_fallback is not None and pos_fallback < len(r_cells):
                return r_cells[pos_fallback]
            return None

        for idx, row in enumerate(rows[1:], start=2):
            if not row or not any(row):
                continue
            try:
                rec_name = get_val(row, "recipient_name", 0)
                phone = get_val(row, "phone", 1)
                email = get_val(row, "email", 2)
                street = get_val(row, "street", 3)
                city = get_val(row, "city", 4)
                state = get_val(row, "state", 5)
                zip_code = get_val(row, "zip_code", 6)
                country = get_val(row, "country", 7)
                sku = get_val(row, "sku", 8)
                qty = get_val(row, "quantity", 9)
                currency = get_val(row, "currency", 10)
                shipping_fee = get_val(row, "shipping_fee", 11)
                tax_amount = get_val(row, "tax_amount", 12)
                discount_amount = get_val(row, "discount_amount", 13)
                other_charges = get_val(row, "other_charges", 14)
                payment_method = get_val(row, "payment_method", 15)
                payment_channel = get_val(row, "payment_channel", 16)
                delivery_method = get_val(row, "delivery_method", 17)
                courier_company = get_val(row, "courier_company", 18)
                tracking_id = get_val(row, "tracking_id", 19)
                customer_note = get_val(row, "customer_note", 20)
                p_stat = get_val(row, "payment_status", 21)
                f_stat = get_val(row, "fulfillment_status", 22)

                rec_name_str = str(rec_name).strip() if rec_name is not None else ""
                street_str = str(street).strip() if street is not None else ""
                city_str = str(city).strip() if city is not None else ""
                sku_str = str(sku).strip() if sku is not None else ""

                if not rec_name_str or not street_str or not city_str or not sku_str:
                    errors.append(f"Row {idx}: Missing required fields (Recipient Name, Street, City, SKU)")
                    continue

                variant = db.query(ProductVariant).filter(ProductVariant.sku == sku_str).first()
                if not variant:
                    errors.append(f"Row {idx}: Product SKU '{sku_str}' not found")
                    continue

                try:
                    qty_val = int(float(qty)) if qty is not None and str(qty).strip() != "" else 1
                except (ValueError, TypeError):
                    qty_val = 1

                try:
                    ship_val = float(shipping_fee) if shipping_fee is not None and str(shipping_fee).strip() != "" else 0.0
                except (ValueError, TypeError):
                    ship_val = 0.0

                try:
                    tax_val = float(tax_amount) if tax_amount is not None and str(tax_amount).strip() != "" else 0.0
                except (ValueError, TypeError):
                    tax_val = 0.0

                try:
                    disc_val = float(discount_amount) if discount_amount is not None and str(discount_amount).strip() != "" else 0.0
                except (ValueError, TypeError):
                    disc_val = 0.0

                try:
                    other_val = float(other_charges) if other_charges is not None and str(other_charges).strip() != "" else 0.0
                except (ValueError, TypeError):
                    other_val = 0.0

                order_create = OrderCreate(
                    business_id=business_id,
                    user_id=uuid.uuid4(),
                    guest_email=str(email).strip() if email is not None and str(email).strip() else None,
                    currency=str(currency).strip() if currency is not None and str(currency).strip() else "BDT",
                    items=[{"variant_id": variant.id, "quantity": qty_val}],
                    shipping_amount=ship_val,
                    tax_amount=tax_val,
                    discount_amount=disc_val,
                    other_charges=other_val,
                    customer_note=str(customer_note).strip() if customer_note is not None and str(customer_note).strip() else None,
                    payment_method=str(payment_method).strip() if payment_method is not None and str(payment_method).strip() else None,
                    payment_channel=str(payment_channel).strip() if payment_channel is not None and str(payment_channel).strip() else None,
                    delivery_method=str(delivery_method).strip() if delivery_method is not None and str(delivery_method).strip() else None,
                    courier_company=str(courier_company).strip() if courier_company is not None and str(courier_company).strip() else None,
                    tracking_id=str(tracking_id).strip() if tracking_id is not None and str(tracking_id).strip() else None,
                    shipping_address={
                        "address_type": "shipping",
                        "recipient_name": rec_name_str,
                        "phone": str(phone).strip() if phone is not None and str(phone).strip() else None,
                        "street": street_str,
                        "city": city_str,
                        "state": str(state).strip() if state is not None and str(state).strip() else None,
                        "zip_code": str(zip_code).strip() if zip_code is not None and str(zip_code).strip() else None,
                        "country": str(country).strip() if country is not None and str(country).strip() else "Bangladesh",
                    }
                )
                created_order = self.create_order(db, order_create)

                p_str = str(p_stat).strip().lower() if p_stat is not None else ""
                f_str = str(f_stat).strip().lower() if f_stat is not None else ""

                p_enum = OrderPaymentStatus(p_str) if p_str in [s.value for s in OrderPaymentStatus] else None
                f_enum = OrderFulfillmentStatus(f_str) if f_str in [s.value for s in OrderFulfillmentStatus] else None

                if p_enum or f_enum:
                    self.update_order_status(
                        db,
                        order_id=created_order.id,
                        update_data=OrderStatusUpdate(payment_status=p_enum, fulfillment_status=f_enum, note="Imported via Excel"),
                        business_id=business_id
                    )

                imported_count += 1
            except Exception as e:
                errors.append(f"Row {idx}: {str(e)}")

        return {"imported_count": imported_count, "errors": errors}
