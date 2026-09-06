import re
import uuid
from io import BytesIO

import openpyxl
from fastapi import HTTPException, status
from sqlalchemy.orm import Session

from app.modules.ecommerce.brands.models import Brand, ProductModel
from app.modules.ecommerce.brands.repository import BrandRepository
from app.modules.ecommerce.brands.schemas import (
    BrandCreate,
    BrandDropdownItem,
    BrandUpdate,
    ProductModelCreate,
    ProductModelUpdate,
)


class BrandService:

    def __init__(self, repository: BrandRepository | None = None):
        self.repository = repository or BrandRepository()

    # --- Brand Logic ---
    def get_brands(
        self,
        db: Session,
        skip: int = 0,
        limit: int = 100,
        business_id: int | None = None,
        is_active: bool | None = None,
    ) -> list[Brand]:
        return self.repository.get_all(
            db,
            skip=skip,
            limit=limit,
            business_id=business_id,
            is_active=is_active,
        )

    def get_brand_dropdown_items(
        self,
        db: Session,
        business_id: int | None = None,
    ) -> list[BrandDropdownItem]:
        brands = self.get_brands(db, business_id=business_id, is_active=True)
        return [
            BrandDropdownItem(
                id=b.id,
                name=b.name,
                slug=b.slug,
                logo_url=b.logo_url,
            )
            for b in brands
        ]

    def get_brand(self, db: Session, brand_id: uuid.UUID) -> Brand:
        brand = self.repository.get_by_id(db, brand_id)
        if not brand:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Brand not found",
            )
        return brand

    def create_brand(self, db: Session, data: BrandCreate) -> Brand:
        if self.repository.get_by_slug(db, data.slug):
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Brand with slug '{data.slug}' already exists",
            )
        return self.repository.create(db, data)

    def update_brand(
        self, db: Session, brand_id: uuid.UUID, data: BrandUpdate
    ) -> Brand:
        brand = self.get_brand(db, brand_id)
        if data.slug is not None and data.slug != brand.slug:
            existing = self.repository.get_by_slug(db, data.slug)
            if existing and existing.id != brand_id:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"Brand with slug '{data.slug}' already exists",
                )
        return self.repository.update(db, brand, data)

    def delete_brand(self, db: Session, brand_id: uuid.UUID) -> None:
        brand = self.get_brand(db, brand_id)
        self.repository.delete(db, brand)

    # --- ProductModel Logic ---
    def get_all_models(
        self,
        db: Session,
        skip: int = 0,
        limit: int = 500,
        is_active: bool | None = None,
    ) -> list[ProductModel]:
        return self.repository.get_all_models(
            db, skip=skip, limit=limit, is_active=is_active
        )

    def get_brand_models(
        self,
        db: Session,
        brand_id: uuid.UUID,
        is_active: bool | None = None,
    ) -> list[ProductModel]:
        self.get_brand(db, brand_id)
        return self.repository.get_models_by_brand(
            db, brand_id, is_active=is_active
        )

    def create_model(
        self, db: Session, brand_id: uuid.UUID, data: ProductModelCreate
    ) -> ProductModel:
        self.get_brand(db, brand_id)
        if self.repository.get_model_by_slug(db, data.slug):
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Model with slug '{data.slug}' already exists",
            )
        return self.repository.create_model(db, brand_id, data)

    def update_model(
        self, db: Session, model_id: uuid.UUID, data: ProductModelUpdate
    ) -> ProductModel:
        model = self.repository.get_model_by_id(db, model_id)
        if not model:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Product model not found",
            )
        if data.slug is not None and data.slug != model.slug:
            existing = self.repository.get_model_by_slug(db, data.slug)
            if existing and existing.id != model_id:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"Model with slug '{data.slug}' already exists",
                )
        return self.repository.update_model(db, model, data)

    def delete_model(self, db: Session, model_id: uuid.UUID) -> None:
        model = self.repository.get_model_by_id(db, model_id)
        if not model:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Product model not found",
            )
        self.repository.delete_model(db, model)

    def generate_excel_template(self, db: Session, business_id: int) -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Brand Template"

        headers = [
            "Brand Name",
            "Slug",
            "Logo URL",
            "Description",
        ]
        ws.append(headers)

        ws.append([
            "Apple",
            "apple",
            "https://example.com/logos/apple.png",
            "Apple Inc. consumer electronics brand",
        ])
        ws.append([
            "Samsung",
            "samsung",
            "https://example.com/logos/samsung.png",
            "Samsung Electronics brand",
        ])

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

        output = BytesIO()
        wb.save(output)
        return output.getvalue()

    def import_brands_excel(
        self, db: Session, business_id: int, file_bytes: bytes
    ) -> dict[str, int | list[str]]:
        try:
            wb = openpyxl.load_workbook(filename=BytesIO(file_bytes), data_only=True)
            ws = wb.active
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Invalid Excel file format: {str(e)}",
            )

        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Excel file is empty.",
            )

        success_count = 0
        error_messages: list[str] = []

        header = [str(cell or "").strip().lower() for cell in rows[0]]
        start_idx = 1 if "brand name" in header or "name" in header or "slug" in header else 0

        for row_idx, row in enumerate(rows[start_idx:], start=start_idx + 1):
            if not row or not any(row):
                continue

            name_raw = str(row[0] or "").strip()
            slug_raw = str(row[1] or "").strip().lower() if len(row) > 1 and row[1] else ""
            logo_raw = str(row[2] or "").strip() if len(row) > 2 and row[2] else None
            desc_raw = str(row[3] or "").strip() if len(row) > 3 and row[3] else None

            if not name_raw:
                error_messages.append(f"Row {row_idx}: Missing Brand Name.")
                continue

            if not slug_raw:
                slug_raw = re.sub(r"[^\w\s-]", "", name_raw).strip().lower().replace(" ", "-")

            existing_name = self.repository.get_by_name(db, name_raw, business_id)
            existing_slug = self.repository.get_by_slug(db, slug_raw)

            if existing_name or existing_slug:
                warning_reason = []
                if existing_name:
                    warning_reason.append(f"Brand name '{name_raw}' already exists")
                if existing_slug:
                    warning_reason.append(f"Slug '{slug_raw}' already exists")
                error_messages.append(f"Row {row_idx}: Skipped - {', '.join(warning_reason)}.")
                continue

            create_data = BrandCreate(
                business_id=business_id,
                name=name_raw,
                slug=slug_raw,
                logo_url=logo_raw,
                description=desc_raw,
                is_active=True,
            )

            try:
                self.create_brand(db, create_data)
                success_count += 1
            except HTTPException as hexp:
                error_messages.append(f"Row {row_idx}: {hexp.detail}")
            except Exception as ex:
                error_messages.append(f"Row {row_idx}: Failed to create brand - {str(ex)}")

        return {
            "imported_count": success_count,
            "errors": error_messages,
        }

    def generate_model_excel_template(self, db: Session, business_id: int) -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Model Template"

        headers = [
            "Model Name",
            "Brand Name",
            "Slug",
            "Description",
        ]
        ws.append(headers)

        brands = self.repository.get_all(db, business_id=business_id, limit=10)
        sample_brand = brands[0].name if brands else "Apple"

        ws.append([
            "iPhone 15 Pro",
            sample_brand,
            "iphone-15-pro",
            "Flagship smartphone series model",
        ])
        ws.append([
            "Galaxy S24 Ultra",
            "Samsung",
            "galaxy-s24-ultra",
            "Premium smartphone series model",
        ])

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

        output = BytesIO()
        wb.save(output)
        return output.getvalue()

    def import_models_excel(
        self, db: Session, business_id: int, file_bytes: bytes
    ) -> dict[str, int | list[str]]:
        try:
            wb = openpyxl.load_workbook(filename=BytesIO(file_bytes), data_only=True)
            ws = wb.active
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Invalid Excel file format: {str(e)}",
            )

        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Excel file is empty.",
            )

        success_count = 0
        error_messages: list[str] = []

        header = [str(cell or "").strip().lower() for cell in rows[0]]
        start_idx = 1 if "model name" in header or "name" in header or "brand name" in header else 0

        for row_idx, row in enumerate(rows[start_idx:], start=start_idx + 1):
            if not row or not any(row):
                continue

            name_raw = str(row[0] or "").strip()
            brand_name_raw = str(row[1] or "").strip() if len(row) > 1 and row[1] else ""
            slug_raw = str(row[2] or "").strip().lower() if len(row) > 2 and row[2] else ""
            desc_raw = str(row[3] or "").strip() if len(row) > 3 and row[3] else None

            if not name_raw:
                error_messages.append(f"Row {row_idx}: Missing Model Name.")
                continue

            if not brand_name_raw:
                error_messages.append(f"Row {row_idx}: Missing Brand Name.")
                continue

            if not slug_raw:
                slug_raw = re.sub(r"[^\w\s-]", "", name_raw).strip().lower().replace(" ", "-")

            brand_obj = self.repository.get_by_name(db, brand_name_raw, business_id)
            if not brand_obj:
                brand_obj = self.repository.get_by_name(db, brand_name_raw, None)

            if not brand_obj:
                error_messages.append(f"Row {row_idx}: Brand '{brand_name_raw}' not found.")
                continue

            existing_model = self.repository.get_model_by_name(db, name_raw, brand_obj.id)
            existing_slug = self.repository.get_model_by_slug(db, slug_raw)

            if existing_model or existing_slug:
                warning_reason = []
                if existing_model:
                    warning_reason.append(f"Model name '{name_raw}' already exists for brand '{brand_name_raw}'")
                if existing_slug:
                    warning_reason.append(f"Slug '{slug_raw}' already exists")
                error_messages.append(f"Row {row_idx}: Skipped - {', '.join(warning_reason)}.")
                continue

            create_data = ProductModelCreate(
                name=name_raw,
                slug=slug_raw,
                description=desc_raw,
                is_active=True,
            )

            try:
                self.create_model(db, brand_obj.id, create_data)
                success_count += 1
            except HTTPException as hexp:
                error_messages.append(f"Row {row_idx}: {hexp.detail}")
            except Exception as ex:
                error_messages.append(f"Row {row_idx}: Failed to create model - {str(ex)}")

        return {
            "imported_count": success_count,
            "errors": error_messages,
        }
