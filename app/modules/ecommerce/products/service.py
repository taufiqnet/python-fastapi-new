import os
import re
import uuid
from decimal import Decimal
from io import BytesIO

import openpyxl
from fastapi import HTTPException, UploadFile, status
from sqlalchemy.orm import Session

from app.common.enums import Status
from app.modules.ecommerce.brands.repository import BrandRepository
from app.modules.ecommerce.categories.repository import CategoryRepository
from app.modules.ecommerce.products.models import (
    Product,
    ProductCondition,
    ProductImage,
    ProductTag,
    ProductType,
    ProductVariant,
)
from app.modules.ecommerce.products.repository import ProductRepository
from app.modules.ecommerce.products.schemas import (
    ProductCreate,
    ProductImageCreate,
    ProductListItem,
    ProductTagCreate,
    ProductUpdate,
    VariantCreate,
    VariantUpdate,
)

MAX_FILE_SIZE = 5 * 1024 * 1024  # 5MB
ALLOWED_MIME_TYPES = {"image/jpeg", "image/png", "image/webp", "image/gif"}


class ProductService:

    def __init__(
        self,
        repository: ProductRepository | None = None,
        category_repository: CategoryRepository | None = None,
        brand_repository: BrandRepository | None = None,
    ):
        self.repository = repository or ProductRepository()
        self.category_repository = category_repository or CategoryRepository()
        self.brand_repository = brand_repository or BrandRepository()

    def get_products(
        self,
        db: Session,
        skip: int = 0,
        limit: int = 100,
        business_id: int | None = None,
        category_id: uuid.UUID | None = None,
        brand_id: uuid.UUID | None = None,
        model_id: uuid.UUID | None = None,
        status: Status | None = None,
    ) -> list[Product]:
        return self.repository.get_all(
            db,
            skip=skip,
            limit=limit,
            business_id=business_id,
            category_id=category_id,
            brand_id=brand_id,
            model_id=model_id,
            status=status,
        )

    def get_product_list_items(
        self,
        db: Session,
        skip: int = 0,
        limit: int = 100,
        business_id: int | None = None,
        category_id: uuid.UUID | None = None,
        brand_id: uuid.UUID | None = None,
        model_id: uuid.UUID | None = None,
        status: Status | None = None,
    ) -> list[ProductListItem]:
        products = self.get_products(
            db,
            skip=skip,
            limit=limit,
            business_id=business_id,
            category_id=category_id,
            brand_id=brand_id,
            model_id=model_id,
            status=status,
        )
        items = []
        for p in products:
            primary_img = next((img.url for img in p.images if img.is_primary), None)
            if not primary_img and p.images:
                primary_img = p.images[0].url

            min_price = None
            if p.variants:
                min_price = min(v.price for v in p.variants)

            items.append(
                ProductListItem(
                    id=p.id,
                    title=p.title,
                    slug=p.slug,
                    brand=p.brand or (p.brand_rel.name if p.brand_rel else None),
                    brand_id=p.brand_id,
                    model_id=p.model_id,
                    status=p.status,
                    thumbnail_url=primary_img,
                    min_price=min_price,
                    business_id=p.business_id,
                    is_featured=p.is_featured,
                    average_rating=p.average_rating,
                )
            )
        return items

    def get_product(self, db: Session, product_id: uuid.UUID) -> Product:
        product = self.repository.get_by_id(db, product_id)
        if not product:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Product not found",
            )
        return product

    def get_product_by_slug(self, db: Session, slug: str) -> Product:
        product = self.repository.get_by_slug(db, slug)
        if not product:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Product not found",
            )
        return product

    def create_product(self, db: Session, data: ProductCreate) -> Product:
        if self.repository.get_by_slug(db, data.slug):
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Product with slug '{data.slug}' already exists",
            )
        return self.repository.create(db, data)

    def update_product(
        self, db: Session, product_id: uuid.UUID, data: ProductUpdate
    ) -> Product:
        product = self.get_product(db, product_id)

        if data.slug is not None and data.slug != product.slug:
            existing = self.repository.get_by_slug(db, data.slug)
            if existing and existing.id != product_id:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"Product with slug '{data.slug}' already exists",
                )

        return self.repository.update(db, product, data)

    def delete_product(self, db: Session, product_id: uuid.UUID) -> None:
        product = self.get_product(db, product_id)

        # Delete physical images
        for img in product.images:
            self._delete_physical_image_file(img.url)

        self.repository.delete(db, product)

    # --- Tags ---
    def create_tag(self, db: Session, data: ProductTagCreate) -> ProductTag:
        if self.repository.get_tag_by_slug(db, data.slug):
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Tag with slug '{data.slug}' already exists",
            )
        return self.repository.create_tag(db, data)

    def get_tags(
        self, db: Session, business_id: int | None = None
    ) -> list[ProductTag]:
        return self.repository.get_tags(db, business_id=business_id)

    # --- Variants ---
    def add_variant(
        self, db: Session, product_id: uuid.UUID, data: VariantCreate
    ) -> ProductVariant:
        self.get_product(db, product_id)
        return self.repository.create_variant(db, product_id, data)

    def update_variant(
        self, db: Session, variant_id: uuid.UUID, data: VariantUpdate
    ) -> ProductVariant:
        variant = self.repository.get_variant_by_id(db, variant_id)
        if not variant:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Variant not found",
            )
        return self.repository.update_variant(db, variant, data)

    def delete_variant(self, db: Session, variant_id: uuid.UUID) -> None:
        variant = self.repository.get_variant_by_id(db, variant_id)
        if not variant:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Variant not found",
            )
        self.repository.delete_variant(db, variant)

    # --- Images ---
    def upload_product_images(
        self, product_slug: str, files: list[UploadFile]
    ) -> list[str]:
        if not files:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="No files provided",
            )

        upload_dir = "app/static/ecommerce/images"
        os.makedirs(upload_dir, exist_ok=True)

        saved_urls = []
        for index, file in enumerate(files):
            if not file.content_type or not file.content_type.startswith("image/"):
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"File '{file.filename}' is not a valid image.",
                )

            # Check size
            file.file.seek(0, os.SEEK_END)
            file_size = file.file.tell()
            file.file.seek(0)

            if file_size > MAX_FILE_SIZE:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"File '{file.filename}' exceeds maximum allowed size of 5MB.",
                )

            # Determine extension
            ext = ".jpg"
            if file.filename and "." in file.filename:
                ext = "." + file.filename.rsplit(".", 1)[-1].lower()

            # Format name: product_slug, product_slug_1, product_slug_2, etc.
            if index == 0 and not os.path.exists(os.path.join(upload_dir, f"{product_slug}{ext}")):
                filename = f"{product_slug}{ext}"
            else:
                counter = index if index > 0 else 1
                while os.path.exists(os.path.join(upload_dir, f"{product_slug}_{counter}{ext}")):
                    counter += 1
                filename = f"{product_slug}_{counter}{ext}"

            filepath = os.path.join(upload_dir, filename)
            with open(filepath, "wb") as f:
                f.write(file.file.read())

            relative_url = f"/static/ecommerce/images/{filename}"
            saved_urls.append(relative_url)

        return saved_urls

    def add_image(
        self, db: Session, product_id: uuid.UUID, data: ProductImageCreate
    ) -> ProductImage:
        self.get_product(db, product_id)
        return self.repository.create_image(db, product_id, data)

    def delete_image(self, db: Session, image_id: uuid.UUID) -> None:
        image = self.repository.get_image_by_id(db, image_id)
        if not image:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Image not found",
            )

        self._delete_physical_image_file(image.url)
        self.repository.delete_image(db, image)

    def _delete_physical_image_file(self, url: str) -> None:
        if url and url.startswith("/static/ecommerce/images/"):
            filename = url.replace("/static/ecommerce/images/", "")
            file_path = os.path.join("app/static/ecommerce/images", filename)
            if os.path.exists(file_path):
                try:
                    os.remove(file_path)
                except Exception:
                    pass

    def generate_excel_template(self, db: Session, business_id: int) -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Product Template"

        headers = [
            "Product Title",
            "Slug",
            "Category Name",
            "Brand Name",
            "Model Name",
            "Status (draft/published/archived)",
            "Condition (new/used/refurbished/open_box)",
            "Product Type (physical/digital/service)",
            "Price",
            "SKU",
            "Stock Quantity",
            "Description",
        ]
        ws.append(headers)

        categories = self.category_repository.get_all(db, business_id=business_id, limit=10)
        sample_cat = categories[0].name if categories else "Smartphones"

        brands = self.brand_repository.get_all(db, business_id=business_id, limit=10)
        sample_brand = brands[0].name if brands else "Apple"

        ws.append([
            "iPhone 15 Pro 128GB",
            "iphone-15-pro-128gb",
            sample_cat,
            sample_brand,
            "",
            "published",
            "new",
            "physical",
            999.99,
            "IPH15P-128-BLK",
            50,
            "Apple iPhone 15 Pro with Titanium design and A17 Pro chip.",
        ])

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

        output = BytesIO()
        wb.save(output)
        return output.getvalue()

    def import_products_excel(
        self, db: Session, business_id: int, file_bytes: bytes
    ) -> dict[str, int | list[str]]:
        try:
            wb = openpyxl.load_workbook(filename=BytesIO(file_bytes), data_only=True)
            ws = wb.active
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Invalid Excel file format: {str(e)}",
            )

        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Excel file is empty.",
            )

        success_count = 0
        error_messages: list[str] = []

        header = [str(cell or "").strip().lower() for cell in rows[0]]
        start_idx = 1 if "product title" in header or "title" in header or "slug" in header else 0

        for row_idx, row in enumerate(rows[start_idx:], start=start_idx + 1):
            if not row or not any(row):
                continue

            title_raw = str(row[0] or "").strip()
            slug_raw = str(row[1] or "").strip().lower() if len(row) > 1 and row[1] else ""
            cat_name_raw = str(row[2] or "").strip() if len(row) > 2 and row[2] else None
            brand_name_raw = str(row[3] or "").strip() if len(row) > 3 and row[3] else None
            model_name_raw = str(row[4] or "").strip() if len(row) > 4 and row[4] else None
            status_raw = str(row[5] or "").strip().lower() if len(row) > 5 and row[5] else "draft"
            condition_raw = str(row[6] or "").strip().lower() if len(row) > 6 and row[6] else "new"
            product_type_raw = str(row[7] or "").strip().lower() if len(row) > 7 and row[7] else "physical"
            price_raw = row[8] if len(row) > 8 else None
            sku_raw = str(row[9] or "").strip() if len(row) > 9 and row[9] else None
            stock_raw = row[10] if len(row) > 10 else None
            desc_raw = str(row[11] or "").strip() if len(row) > 11 and row[11] else None

            if not title_raw:
                error_messages.append(f"Row {row_idx}: Missing Product Title.")
                continue

            if not slug_raw:
                slug_raw = re.sub(r"[^\w\s-]", "", title_raw).strip().lower().replace(" ", "-")

            existing_title = self.repository.get_by_title(db, title_raw, business_id)
            existing_slug = self.repository.get_by_slug(db, slug_raw)

            if existing_title or existing_slug:
                warning_reason = []
                if existing_title:
                    warning_reason.append(f"Product title '{title_raw}' already exists")
                if existing_slug:
                    warning_reason.append(f"Slug '{slug_raw}' already exists")
                error_messages.append(f"Row {row_idx}: Skipped - {', '.join(warning_reason)}.")
                continue

            cat_id = None
            if cat_name_raw:
                cat_obj = self.category_repository.get_by_name(db, cat_name_raw, business_id)
                if not cat_obj:
                    cat_obj = self.category_repository.get_by_name(db, cat_name_raw, None)
                if cat_obj:
                    cat_id = cat_obj.id
                else:
                    error_messages.append(f"Row {row_idx}: Category '{cat_name_raw}' not found.")
                    continue

            brand_id = None
            if brand_name_raw:
                brand_obj = self.brand_repository.get_by_name(db, brand_name_raw, business_id)
                if not brand_obj:
                    brand_obj = self.brand_repository.get_by_name(db, brand_name_raw, None)
                if brand_obj:
                    brand_id = brand_obj.id
                else:
                    error_messages.append(f"Row {row_idx}: Brand '{brand_name_raw}' not found.")
                    continue

            model_id = None
            if model_name_raw:
                model_obj = self.brand_repository.get_model_by_name(db, model_name_raw, brand_id)
                if model_obj:
                    model_id = model_obj.id
                else:
                    error_messages.append(f"Row {row_idx}: Model '{model_name_raw}' not found.")
                    continue

            # Parse Status
            st_enum = Status.DRAFT
            if status_raw:
                try:
                    st_enum = Status(status_raw)
                except ValueError:
                    pass

            # Parse Condition
            cond_enum = ProductCondition.NEW
            if condition_raw:
                try:
                    cond_enum = ProductCondition(condition_raw)
                except ValueError:
                    pass

            # Parse Product Type
            pt_enum = ProductType.PHYSICAL
            if product_type_raw:
                try:
                    pt_enum = ProductType(product_type_raw)
                except ValueError:
                    pass

            # Parse Variant
            variants: list[VariantCreate] = []
            if price_raw is not None or sku_raw is not None:
                try:
                    parsed_price = Decimal(str(price_raw or "0.00").strip())
                except Exception:
                    parsed_price = Decimal("0.00")

                if parsed_price <= Decimal("0.00"):
                    parsed_price = Decimal("1.00")

                if not sku_raw:
                    sku_raw = f"{slug_raw.upper()[:30]}-DEFAULT"

                existing_v = self.repository.get_variant_by_sku(db, sku_raw)
                if existing_v:
                    error_messages.append(f"Row {row_idx}: Skipped - Variant SKU '{sku_raw}' already exists.")
                    continue

                try:
                    parsed_stock = int(stock_raw) if stock_raw is not None else 0
                except Exception:
                    parsed_stock = 0

                variants.append(
                    VariantCreate(
                        sku=sku_raw,
                        price=parsed_price,
                        stock_qty=parsed_stock,
                        is_default=True,
                    )
                )

            create_data = ProductCreate(
                business_id=business_id,
                title=title_raw,
                slug=slug_raw,
                category_id=cat_id,
                brand_id=brand_id,
                model_id=model_id,
                brand=brand_name_raw,
                status=st_enum,
                condition=cond_enum,
                product_type=pt_enum,
                description=desc_raw,
                variants=variants,
            )

            try:
                self.create_product(db, create_data)
                success_count += 1
            except HTTPException as hexp:
                error_messages.append(f"Row {row_idx}: {hexp.detail}")
            except Exception as ex:
                error_messages.append(f"Row {row_idx}: Failed to create product - {str(ex)}")

        return {
            "imported_count": success_count,
            "errors": error_messages,
        }
