import uuid
from io import BytesIO

import openpyxl
from fastapi import HTTPException, status
from sqlalchemy.orm import Session

from app.modules.ecommerce.customer.models import Customer
from app.modules.ecommerce.customer.repository import CustomerRepository
from app.modules.ecommerce.customer.schemas import CustomerCreate, CustomerUpdate


class CustomerService:

    def __init__(self, repository: CustomerRepository | None = None):
        self.repository = repository or CustomerRepository()

    def get_customers(
        self,
        db: Session,
        skip: int = 0,
        limit: int = 100,
        business_id: int | None = None,
        is_active: bool | None = None,
    ) -> list[Customer]:
        return self.repository.get_all(
            db, skip=skip, limit=limit, business_id=business_id, is_active=is_active
        )

    def get_customer(self, db: Session, customer_id: uuid.UUID) -> Customer:
        customer = self.repository.get_by_id(db, customer_id)
        if not customer:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Customer not found",
            )
        return customer

    def create_customer(self, db: Session, data: CustomerCreate) -> Customer:
        if not data.first_name or not data.first_name.strip():
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="First name is required",
            )
        if not data.phone or not data.phone.strip():
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Phone number is required",
            )
        if not data.address or not data.address.strip():
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Full Address is required",
            )

        existing_customer = self.repository.get_by_phone(
            db, data.phone.strip(), business_id=data.business_id
        )
        if existing_customer:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="A customer with this phone number already exists.",
            )

        return self.repository.create(db, data)

    def update_customer(
        self, db: Session, customer_id: uuid.UUID, data: CustomerUpdate
    ) -> Customer:
        customer = self.get_customer(db, customer_id)

        if data.first_name is not None and not data.first_name.strip():
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="First name cannot be empty",
            )
        if data.phone is not None and not data.phone.strip():
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Phone number cannot be empty",
            )
        if data.address is not None and not data.address.strip():
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Full Address cannot be empty",
            )

        return self.repository.update(db, customer, data)

    def delete_customer(self, db: Session, customer_id: uuid.UUID) -> None:
        customer = self.get_customer(db, customer_id)
        self.repository.delete(db, customer)

    def generate_excel_template(self, db: Session, business_id: int) -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Customer Template"

        headers = [
            "First Name",
            "Last Name",
            "Phone",
            "Email",
            "Address",
        ]
        ws.append(headers)

        ws.append([
            "John",
            "Doe",
            "+8801700112233",
            "john.doe@example.com",
            "123 Main Street, Dhaka",
        ])
        ws.append([
            "Jane",
            "Smith",
            "+8801800112233",
            "jane.smith@example.com",
            "456 Oak Avenue, Chittagong",
        ])

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

        output = BytesIO()
        wb.save(output)
        return output.getvalue()

    def import_customers_excel(
        self, db: Session, business_id: int, file_bytes: bytes
    ) -> dict[str, int | list[str]]:
        try:
            wb = openpyxl.load_workbook(filename=BytesIO(file_bytes), data_only=True)
            ws = wb.active
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Invalid Excel file format: {str(e)}",
            )

        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Excel file is empty.",
            )

        success_count = 0
        error_messages: list[str] = []

        header = [str(cell or "").strip().lower() for cell in rows[0]]

        first_name_idx = 0
        last_name_idx = 1
        phone_idx = 2
        email_idx = 3
        address_idx = 4

        has_header = False
        for idx, col_name in enumerate(header):
            if "first" in col_name or ("name" in col_name and "last" not in col_name):
                first_name_idx = idx
                has_header = True
            elif "last" in col_name:
                last_name_idx = idx
                has_header = True
            elif "phone" in col_name or "mobile" in col_name:
                phone_idx = idx
                has_header = True
            elif "email" in col_name:
                email_idx = idx
                has_header = True
            elif "address" in col_name:
                address_idx = idx
                has_header = True

        start_idx = 1 if has_header else 0

        for row_idx, row in enumerate(rows[start_idx:], start=start_idx + 1):
            if not row or not any(row):
                continue

            first_name_raw = (
                str(row[first_name_idx] or "").strip()
                if len(row) > first_name_idx and row[first_name_idx] is not None
                else ""
            )
            last_name_raw = (
                str(row[last_name_idx] or "").strip()
                if len(row) > last_name_idx and row[last_name_idx] is not None
                else None
            )
            if last_name_raw == "":
                last_name_raw = None

            phone_raw = (
                str(row[phone_idx] or "").strip()
                if len(row) > phone_idx and row[phone_idx] is not None
                else ""
            )
            email_raw = (
                str(row[email_idx] or "").strip()
                if len(row) > email_idx and row[email_idx] is not None
                else None
            )
            if email_raw == "":
                email_raw = None

            address_raw = (
                str(row[address_idx] or "").strip()
                if len(row) > address_idx and row[address_idx] is not None
                else ""
            )

            if not first_name_raw:
                error_messages.append(f"Row {row_idx}: Missing First Name.")
                continue

            if not phone_raw:
                error_messages.append(f"Row {row_idx}: Missing Phone Number.")
                continue

            if not address_raw:
                error_messages.append(f"Row {row_idx}: Missing Address.")
                continue

            existing_phone = self.repository.get_by_phone(
                db, phone_raw, business_id=business_id
            )
            existing_email = (
                self.repository.get_by_email(db, email_raw, business_id=business_id)
                if email_raw
                else None
            )

            if existing_phone or existing_email:
                warning_reason = []
                if existing_phone:
                    warning_reason.append(f"Phone '{phone_raw}' already exists")
                if existing_email:
                    warning_reason.append(f"Email '{email_raw}' already exists")
                error_messages.append(
                    f"Row {row_idx}: Skipped - Customer already exists ({', '.join(warning_reason)})."
                )
                continue

            create_data = CustomerCreate(
                first_name=first_name_raw,
                last_name=last_name_raw,
                phone=phone_raw,
                email=email_raw,
                address=address_raw,
                business_id=business_id,
                is_active=True,
            )

            try:
                self.create_customer(db, create_data)
                success_count += 1
            except HTTPException as hexp:
                error_messages.append(f"Row {row_idx}: {hexp.detail}")
            except Exception as ex:
                error_messages.append(
                    f"Row {row_idx}: Failed to create customer - {str(ex)}"
                )

        return {
            "imported_count": success_count,
            "errors": error_messages,
        }
