import re
import uuid
from io import BytesIO

import openpyxl
from fastapi import HTTPException, status
from sqlalchemy.orm import Session

from app.modules.ecommerce.categories.models import Category
from app.modules.ecommerce.categories.repository import CategoryRepository
from app.modules.ecommerce.categories.schemas import CategoryCreate, CategoryUpdate


class CategoryService:

    def __init__(self, repository: CategoryRepository | None = None):
        self.repository = repository or CategoryRepository()

    def get_categories(
        self,
        db: Session,
        skip: int = 0,
        limit: int = 100,
        business_id: int | None = None,
        parent_id: uuid.UUID | None = None,
    ) -> list[Category]:
        return self.repository.get_all(
            db, skip=skip, limit=limit, business_id=business_id, parent_id=parent_id
        )

    def get_category_tree(
        self, db: Session, business_id: int | None = None
    ) -> list[Category]:
        return self.repository.get_tree(db, business_id=business_id)

    def get_category(self, db: Session, category_id: uuid.UUID) -> Category:
        category = self.repository.get_by_id(db, category_id)
        if not category:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Category not found",
            )
        return category

    def get_category_by_slug(self, db: Session, slug: str) -> Category:
        category = self.repository.get_by_slug(db, slug)
        if not category:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Category not found",
            )
        return category

    def create_category(self, db: Session, data: CategoryCreate) -> Category:
        if self.repository.get_by_slug(db, data.slug):
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Category with slug '{data.slug}' already exists",
            )

        if data.parent_id is not None:
            parent = self.repository.get_by_id(db, data.parent_id)
            if not parent:
                raise HTTPException(
                    status_code=status.HTTP_404_NOT_FOUND,
                    detail=f"Parent category with id '{data.parent_id}' not found",
                )

        return self.repository.create(db, data)

    def update_category(
        self, db: Session, category_id: uuid.UUID, data: CategoryUpdate
    ) -> Category:
        category = self.get_category(db, category_id)

        if data.slug is not None and data.slug != category.slug:
            existing = self.repository.get_by_slug(db, data.slug)
            if existing and existing.id != category_id:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"Category with slug '{data.slug}' already exists",
                )

        if data.parent_id is not None:
            if data.parent_id == category_id:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="Category cannot be its own parent",
                )
            parent = self.repository.get_by_id(db, data.parent_id)
            if not parent:
                raise HTTPException(
                    status_code=status.HTTP_404_NOT_FOUND,
                    detail=f"Parent category with id '{data.parent_id}' not found",
                )

        return self.repository.update(db, category, data)

    def delete_category(self, db: Session, category_id: uuid.UUID) -> None:
        category = self.get_category(db, category_id)
        self.repository.delete(db, category)

    def generate_excel_template(self, db: Session, business_id: int) -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Category Template"

        headers = [
            "Category Name",
            "Slug",
            "Parent Category Name",
            "Description",
            "Icon",
        ]
        ws.append(headers)

        ws.append([
            "Electronics",
            "electronics",
            "",
            "Gadgets and electronic items",
            "fas fa-laptop",
        ])
        ws.append([
            "Smartphones",
            "smartphones",
            "Electronics",
            "Mobile smartphones and accessories",
            "fas fa-mobile-alt",
        ])

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

        output = BytesIO()
        wb.save(output)
        return output.getvalue()

    def import_categories_excel(
        self, db: Session, business_id: int, file_bytes: bytes
    ) -> dict[str, int | list[str]]:
        try:
            wb = openpyxl.load_workbook(filename=BytesIO(file_bytes), data_only=True)
            ws = wb.active
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Invalid Excel file format: {str(e)}",
            )

        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Excel file is empty.",
            )

        success_count = 0
        error_messages: list[str] = []

        header = [str(cell or "").strip().lower() for cell in rows[0]]
        start_idx = 1 if "category name" in header or "name" in header or "slug" in header else 0

        for row_idx, row in enumerate(rows[start_idx:], start=start_idx + 1):
            if not row or not any(row):
                continue

            name_raw = str(row[0] or "").strip()
            slug_raw = str(row[1] or "").strip().lower() if len(row) > 1 and row[1] else ""
            parent_name_raw = str(row[2] or "").strip() if len(row) > 2 and row[2] else None
            desc_raw = str(row[3] or "").strip() if len(row) > 3 and row[3] else None
            icon_raw = str(row[4] or "").strip() if len(row) > 4 and row[4] else None

            if not name_raw:
                error_messages.append(f"Row {row_idx}: Missing Category Name.")
                continue

            if not slug_raw:
                slug_raw = re.sub(r"[^\w\s-]", "", name_raw).strip().lower().replace(" ", "-")

            existing_name = self.repository.get_by_name(db, name_raw, business_id)
            existing_slug = self.repository.get_by_slug(db, slug_raw)

            if existing_name or existing_slug:
                warning_reason = []
                if existing_name:
                    warning_reason.append(f"Category name '{name_raw}' already exists")
                if existing_slug:
                    warning_reason.append(f"Slug '{slug_raw}' already exists")
                error_messages.append(f"Row {row_idx}: Skipped - {', '.join(warning_reason)}.")
                continue

            parent_id = None
            if parent_name_raw:
                parent_cat = self.repository.get_by_name(db, parent_name_raw, business_id)
                if not parent_cat:
                    parent_cat = self.repository.get_by_name(db, parent_name_raw, None)
                if parent_cat:
                    parent_id = parent_cat.id
                else:
                    error_messages.append(f"Row {row_idx}: Parent category '{parent_name_raw}' not found.")
                    continue

            create_data = CategoryCreate(
                business_id=business_id,
                name=name_raw,
                slug=slug_raw,
                parent_id=parent_id,
                description=desc_raw,
                icon=icon_raw,
                is_active=True,
            )

            try:
                self.create_category(db, create_data)
                success_count += 1
            except HTTPException as hexp:
                error_messages.append(f"Row {row_idx}: {hexp.detail}")
            except Exception as ex:
                error_messages.append(f"Row {row_idx}: Failed to create category - {str(ex)}")

        return {
            "imported_count": success_count,
            "errors": error_messages,
        }
