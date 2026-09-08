import uuid

from fastapi import HTTPException, status
from sqlalchemy.orm import Session

from app.modules.hr_payroll.compensation.models import EmployeeSalary
from app.modules.hr_payroll.compensation.repository import EmployeeSalaryRepository
from app.modules.hr_payroll.compensation.schemas import (
    EmployeeSalaryCreate,
    EmployeeSalaryUpdate,
)
from app.modules.hr_payroll.employees.repository import EmployeeRepository
from datetime import date, datetime
from io import BytesIO
import openpyxl


class EmployeeSalaryService:
    def __init__(
        self,
        repository: EmployeeSalaryRepository | None = None,
        employee_repository: EmployeeRepository | None = None,
    ):
        self.repository = repository or EmployeeSalaryRepository()
        self.employee_repository = employee_repository or EmployeeRepository()

    def _compute_totals(
        self,
        basic_salary: float,
        house_rent: float,
        medical_allowance: float,
        transport_allowance: float,
        food_allowance: float,
        other_allowance: float,
        tax: float,
        provident_fund: float,
        other_deduction: float,
    ) -> tuple[float, float]:
        gross = (
            float(basic_salary)
            + float(house_rent)
            + float(medical_allowance)
            + float(transport_allowance)
            + float(food_allowance)
            + float(other_allowance)
        )
        deductions = float(tax) + float(provident_fund) + float(other_deduction)
        net = gross - deductions
        return round(gross, 2), round(net, 2)

    def get_salaries(
        self,
        db: Session,
        skip: int = 0,
        limit: int = 100,
        business_id: int | None = None,
        employee_id: uuid.UUID | None = None,
    ) -> list[EmployeeSalary]:
        return self.repository.get_all(
            db,
            skip=skip,
            limit=limit,
            business_id=business_id,
            employee_id=employee_id,
        )

    def get_salary(self, db: Session, salary_uuid: uuid.UUID) -> EmployeeSalary:
        salary = self.repository.get_by_id(db, salary_uuid)
        if not salary:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Employee salary structure not found",
            )
        return salary

    def get_salary_by_employee(
        self, db: Session, employee_id: uuid.UUID
    ) -> EmployeeSalary:
        salary = self.repository.get_by_employee_id(db, employee_id)
        if not salary:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Salary structure for employee id '{employee_id}' not found",
            )
        return salary

    def create_salary_structure(
        self, db: Session, data: EmployeeSalaryCreate
    ) -> EmployeeSalary:
        employee = self.employee_repository.get_by_id(db, data.employee_id)
        if not employee:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Employee with id '{data.employee_id}' not found",
            )

        existing = self.repository.get_by_employee_id(db, data.employee_id)
        if existing:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=(
                    f"Salary structure already exists for employee "
                    f"'{data.employee_id}'"
                ),
            )

        gross_salary, net_salary = self._compute_totals(
            basic_salary=data.basic_salary,
            house_rent=data.house_rent,
            medical_allowance=data.medical_allowance,
            transport_allowance=data.transport_allowance,
            food_allowance=data.food_allowance,
            other_allowance=data.other_allowance,
            tax=data.tax,
            provident_fund=data.provident_fund,
            other_deduction=data.other_deduction,
        )

        return self.repository.create(
            db, data=data, gross_salary=gross_salary, net_salary=net_salary
        )

    def update_salary_structure(
        self, db: Session, salary_uuid: uuid.UUID, data: EmployeeSalaryUpdate
    ) -> EmployeeSalary:
        salary = self.get_salary(db, salary_uuid)

        basic_salary = (
            data.basic_salary
            if data.basic_salary is not None
            else float(salary.basic_salary)
        )
        house_rent = (
            data.house_rent if data.house_rent is not None else float(salary.house_rent)
        )
        medical_allowance = (
            data.medical_allowance
            if data.medical_allowance is not None
            else float(salary.medical_allowance)
        )
        transport_allowance = (
            data.transport_allowance
            if data.transport_allowance is not None
            else float(salary.transport_allowance)
        )
        food_allowance = (
            data.food_allowance
            if data.food_allowance is not None
            else float(salary.food_allowance)
        )
        other_allowance = (
            data.other_allowance
            if data.other_allowance is not None
            else float(salary.other_allowance)
        )
        tax = data.tax if data.tax is not None else float(salary.tax)
        provident_fund = (
            data.provident_fund
            if data.provident_fund is not None
            else float(salary.provident_fund)
        )
        other_deduction = (
            data.other_deduction
            if data.other_deduction is not None
            else float(salary.other_deduction)
        )

        gross_salary, net_salary = self._compute_totals(
            basic_salary=basic_salary,
            house_rent=house_rent,
            medical_allowance=medical_allowance,
            transport_allowance=transport_allowance,
            food_allowance=food_allowance,
            other_allowance=other_allowance,
            tax=tax,
            provident_fund=provident_fund,
            other_deduction=other_deduction,
        )

        return self.repository.update(
            db,
            salary=salary,
            data=data,
            gross_salary=gross_salary,
            net_salary=net_salary,
        )

    def upsert_salary_structure(
        self, db: Session, data: EmployeeSalaryCreate
    ) -> EmployeeSalary:
        existing = self.repository.get_by_employee_id(db, data.employee_id)
        if existing:
            update_data = EmployeeSalaryUpdate(**data.model_dump(exclude_unset=True))
            return self.update_salary_structure(db, existing.id, update_data)
        return self.create_salary_structure(db, data)

    def delete_salary_structure(self, db: Session, salary_uuid: uuid.UUID) -> None:
        salary = self.get_salary(db, salary_uuid)
        self.repository.delete(db, salary)

    def generate_export_excel(self, db: Session, business_id: int) -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Compensation"

        headers = [
            "Business Profile ID",
            "Business Profile Name",
            "Employee ID",
            "Employee Name",
            "Basic Salary",
            "House Rent",
            "Medical Allowance",
            "Transport Allowance",
            "Food Allowance",
            "Other Allowance",
            "Tax",
            "Provident Fund",
            "Other Deduction",
            "Gross Salary",
            "Net Salary",
            "Effective From",
        ]
        ws.append(headers)

        salaries = self.get_salaries(db, business_id=business_id, limit=2000)
        for s in salaries:
            business_name = s.business_profile.name if s.business_profile else ""
            emp_code = s.employee.employee_id if s.employee else ""
            emp_name = s.employee.full_name if s.employee else ""
            eff_date = s.effective_from.strftime("%Y-%m-%d") if s.effective_from else ""

            ws.append([
                s.business_id,
                business_name,
                emp_code,
                emp_name,
                float(s.basic_salary or 0),
                float(s.house_rent or 0),
                float(s.medical_allowance or 0),
                float(s.transport_allowance or 0),
                float(s.food_allowance or 0),
                float(s.other_allowance or 0),
                float(s.tax or 0),
                float(s.provident_fund or 0),
                float(s.other_deduction or 0),
                float(s.gross_salary or 0),
                float(s.net_salary or 0),
                eff_date,
            ])

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

        output = BytesIO()
        wb.save(output)
        return output.getvalue()

    def generate_excel_template(self, db: Session, business_id: int) -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Compensation Template"

        headers = [
            "Employee ID",
            "Employee Name",
            "Basic Salary",
            "House Rent",
            "Medical Allowance",
            "Transport Allowance",
            "Food Allowance",
            "Other Allowance",
            "Tax",
            "Provident Fund",
            "Other Deduction",
            "Effective From (YYYY-MM-DD)",
        ]
        ws.append(headers)

        employees = self.employee_repository.get_all(db, business_id=business_id, limit=500)
        sample_date = date.today().strftime("%Y-%m-%d")

        if employees:
            for emp in employees[:5]:
                ws.append([
                    emp.employee_id,
                    emp.full_name,
                    5000.0,
                    1500.0,
                    500.0,
                    300.0,
                    200.0,
                    0.0,
                    500.0,
                    250.0,
                    0.0,
                    sample_date,
                ])
        else:
            ws.append([
                "EMP001",
                "John Doe",
                5000.0,
                1500.0,
                500.0,
                300.0,
                200.0,
                0.0,
                500.0,
                250.0,
                0.0,
                sample_date,
            ])

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

        output = BytesIO()
        wb.save(output)
        return output.getvalue()

    def import_salaries_excel(
        self, db: Session, business_id: int, file_bytes: bytes
    ) -> dict[str, int | list[str]]:
        try:
            wb = openpyxl.load_workbook(filename=BytesIO(file_bytes), data_only=True)
            ws = wb.active
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Invalid Excel file format: {str(e)}",
            )

        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Excel file is empty.",
            )

        employees = self.employee_repository.get_all(db, business_id=business_id, limit=2000)
        emp_map = {str(e.employee_id).strip().lower(): e for e in employees}

        success_count = 0
        error_messages: list[str] = []

        header = [str(cell or "").strip().lower() for cell in rows[0]]
        start_idx = 1 if "employee id" in header or "employee_id" in header or "basic salary" in header else 0

        for row_idx, row in enumerate(rows[start_idx:], start=start_idx + 1):
            if not row or not any(row):
                continue

            emp_code_raw = str(row[0] or "").strip()
            basic_sal_raw = row[2] if len(row) > 2 else 0.0
            house_rent_raw = row[3] if len(row) > 3 else 0.0
            medical_raw = row[4] if len(row) > 4 else 0.0
            transport_raw = row[5] if len(row) > 5 else 0.0
            food_raw = row[6] if len(row) > 6 else 0.0
            other_allow_raw = row[7] if len(row) > 7 else 0.0
            tax_raw = row[8] if len(row) > 8 else 0.0
            pf_raw = row[9] if len(row) > 9 else 0.0
            other_ded_raw = row[10] if len(row) > 10 else 0.0
            effective_raw = row[11] if len(row) > 11 else None

            if not emp_code_raw:
                error_messages.append(f"Row {row_idx}: Missing Employee ID.")
                continue

            emp = emp_map.get(emp_code_raw.lower())
            if not emp:
                error_messages.append(f"Row {row_idx}: Employee with ID '{emp_code_raw}' not found.")
                continue

            def parse_float(val) -> float:
                try:
                    return float(val) if val is not None and str(val).strip() != "" else 0.0
                except (ValueError, TypeError):
                    return 0.0

            # Parse effective date
            effective_date: date | None = None
            if isinstance(effective_raw, (datetime, date)):
                effective_date = effective_raw.date() if isinstance(effective_raw, datetime) else effective_raw
            elif isinstance(effective_raw, str) and effective_raw.strip():
                try:
                    effective_date = datetime.strptime(effective_raw.strip(), "%Y-%m-%d").date()
                except ValueError:
                    try:
                        effective_date = datetime.strptime(effective_raw.strip(), "%m/%d/%Y").date()
                    except ValueError:
                        pass

            if not effective_date:
                effective_date = date.today()

            existing = self.repository.get_by_employee_id(db, emp.id)
            if existing:
                error_messages.append(
                    f"Row {row_idx}: Warning - Compensation structure already exists for employee '{emp.full_name}' ({emp_code_raw})."
                )
                continue

            create_data = EmployeeSalaryCreate(
                business_id=business_id,
                employee_id=emp.id,
                basic_salary=parse_float(basic_sal_raw),
                house_rent=parse_float(house_rent_raw),
                medical_allowance=parse_float(medical_raw),
                transport_allowance=parse_float(transport_raw),
                food_allowance=parse_float(food_raw),
                other_allowance=parse_float(other_allow_raw),
                tax=parse_float(tax_raw),
                provident_fund=parse_float(pf_raw),
                other_deduction=parse_float(other_ded_raw),
                effective_from=effective_date,
            )

            try:
                self.create_salary_structure(db, create_data)
                success_count += 1
            except HTTPException as hexp:
                error_messages.append(f"Row {row_idx}: {hexp.detail}")
            except Exception as ex:
                error_messages.append(f"Row {row_idx}: Failed to create compensation structure - {str(ex)}")

        return {
            "imported_count": success_count,
            "errors": error_messages,
        }
