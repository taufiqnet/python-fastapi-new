import uuid

from fastapi import HTTPException, status
from sqlalchemy.orm import Session

from datetime import timedelta

from app.modules.hr_payroll.attendance.models import AttendanceStatusEnum
from app.modules.hr_payroll.attendance.repository import AttendanceRepository
from app.modules.hr_payroll.compensation.repository import EmployeeSalaryRepository
from app.modules.hr_payroll.employees.repository import EmployeeRepository
from app.modules.hr_payroll.leave.models import LeaveStatusEnum
from app.modules.hr_payroll.leave.repository import LeaveApplicationRepository
from app.modules.hr_payroll.payroll.models import (
    Holiday,
    HolidayTypeEnum,
    PaymentMethodEnum,
    PayrollPeriod,
    PayrollPeriodStatusEnum,
    PayrollRecord,
    PayrollSettings,
)
from app.modules.hr_payroll.payroll.repository import (
    HolidayRepository,
    PayrollPeriodRepository,
    PayrollRecordRepository,
    PayrollSettingsRepository,
)
from io import BytesIO
import openpyxl
from app.core.tenancy.repository import BusinessRepository
from app.modules.hr_payroll.payroll.schemas import (
    HolidayCreate,
    HolidayUpdate,
    PayrollPeriodCreate,
    PayrollPeriodUpdate,
    PayrollRecordCreate,
    PayrollRecordUpdate,
    PayrollSettingsUpdate,
)


class HolidayService:
    def __init__(self, repository: HolidayRepository | None = None):
        self.repository = repository or HolidayRepository()

    def get_holidays(
        self,
        db: Session,
        skip: int = 0,
        limit: int = 100,
        business_id: int | None = None,
        holiday_type: str | HolidayTypeEnum | None = None,
    ) -> list[Holiday]:
        return self.repository.get_all(
            db,
            skip=skip,
            limit=limit,
            business_id=business_id,
            holiday_type=holiday_type,
        )

    def get_holiday(self, db: Session, holiday_uuid: uuid.UUID) -> Holiday:
        holiday = self.repository.get_by_id(db, holiday_uuid)
        if not holiday:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Holiday not found",
            )
        return holiday

    def create_holiday(self, db: Session, data: HolidayCreate) -> Holiday:
        if data.start_date > data.end_date:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Start date cannot be after end date",
            )
        return self.repository.create(db, data)

    def update_holiday(
        self, db: Session, holiday_uuid: uuid.UUID, data: HolidayUpdate
    ) -> Holiday:
        holiday = self.get_holiday(db, holiday_uuid)
        start_date = data.start_date or holiday.start_date
        end_date = data.end_date or holiday.end_date
        if start_date > end_date:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Start date cannot be after end date",
            )
        return self.repository.update(db, holiday, data)

    def delete_holiday(self, db: Session, holiday_uuid: uuid.UUID) -> None:
        holiday = self.get_holiday(db, holiday_uuid)
        self.repository.delete(db, holiday)

    def generate_export_excel(self, db: Session, business_id: int | None = None) -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Holidays"

        headers = [
            "Holiday ID",
            "Holiday Name",
            "Holiday Type",
            "Start Date",
            "End Date",
            "Description",
            "Status",
            "Business Profile ID",
            "Business Profile Name",
        ]
        ws.append(headers)

        holidays = self.repository.get_all(db, skip=0, limit=2000, business_id=business_id)
        biz_repo = BusinessRepository()
        businesses = {b.id: b.name_en for b in biz_repo.get_all(db, skip=0, limit=1000)}

        for h in holidays:
            ws.append([
                str(h.id),
                h.name or "",
                h.holiday_type.value if h.holiday_type else "",
                h.start_date.strftime("%Y-%m-%d") if h.start_date else "",
                h.end_date.strftime("%Y-%m-%d") if h.end_date else "",
                h.description or "",
                "Active" if h.is_active else "Inactive",
                str(h.business_id) if h.business_id else "",
                businesses.get(h.business_id, "") if h.business_id else "Global",
            ])

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        output = BytesIO()
        wb.save(output)
        return output.getvalue()


class PayrollPeriodService:
    def __init__(self, repository: PayrollPeriodRepository | None = None):
        self.repository = repository or PayrollPeriodRepository()

    def get_periods(
        self,
        db: Session,
        skip: int = 0,
        limit: int = 100,
        business_id: int | None = None,
        status_filter: str | PayrollPeriodStatusEnum | None = None,
    ) -> list[PayrollPeriod]:
        return self.repository.get_all(
            db,
            skip=skip,
            limit=limit,
            business_id=business_id,
            status=status_filter,
        )

    def get_period(self, db: Session, period_uuid: uuid.UUID) -> PayrollPeriod:
        period = self.repository.get_by_id(db, period_uuid)
        if not period:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Payroll period not found",
            )
        return period

    def create_period(self, db: Session, data: PayrollPeriodCreate) -> PayrollPeriod:
        if data.start_date > data.end_date:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Start date cannot be after end date",
            )

        existing = self.repository.get_by_name(db, data.name, data.business_id)
        if existing:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Payroll period with name '{data.name}' already exists",
            )

        return self.repository.create(db, data)

    def update_period(
        self, db: Session, period_uuid: uuid.UUID, data: PayrollPeriodUpdate
    ) -> PayrollPeriod:
        period = self.get_period(db, period_uuid)

        target_business_id = (
            data.business_id if data.business_id is not None else period.business_id
        )

        if data.name is not None and data.name != period.name:
            existing = self.repository.get_by_name(db, data.name, target_business_id)
            if existing and existing.id != period_uuid:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"Payroll period with name '{data.name}' already exists",
                )

        if period.status == PayrollPeriodStatusEnum.PAID and data.status is not None:
            target_status = (
                data.status
                if isinstance(data.status, PayrollPeriodStatusEnum)
                else PayrollPeriodStatusEnum(data.status)
            )
            if target_status != PayrollPeriodStatusEnum.PAID:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="Cannot modify or revert a paid payroll period",
                )

        start_date = data.start_date or period.start_date
        end_date = data.end_date or period.end_date
        if start_date > end_date:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Start date cannot be after end date",
            )

        return self.repository.update(db, period, data)

    def delete_period(self, db: Session, period_uuid: uuid.UUID) -> None:
        period = self.get_period(db, period_uuid)
        if period.is_locked:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Cannot delete a locked or paid payroll period",
            )
        self.repository.delete(db, period)

    def get_payroll_status(
        self, db: Session, period_uuid: uuid.UUID, business_id: int | None = None
    ) -> dict:
        period = self.get_period(db, period_uuid)
        if business_id is not None and period.business_id != business_id:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Payroll period does not belong to specified business profile",
            )

        record_repo = PayrollRecordRepository()
        records = record_repo.get_all(
            db, business_id=period.business_id, period_id=period_uuid, limit=5000
        )

        total_records = len(records)
        paid_records = sum(1 for r in records if r.is_paid)
        unpaid_records = total_records - paid_records

        return {
            "period_id": str(period.id),
            "period_name": period.name,
            "business_id": period.business_id,
            "start_date": period.start_date.strftime("%Y-%m-%d") if period.start_date else None,
            "end_date": period.end_date.strftime("%Y-%m-%d") if period.end_date else None,
            "status": period.status.value if hasattr(period.status, "value") else str(period.status),
            "is_locked": period.is_locked,
            "total_records": total_records,
            "paid_records": paid_records,
            "unpaid_records": unpaid_records,
        }

    def generate_export_excel(self, db: Session, business_id: int | None = None) -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Payroll Periods"

        headers = [
            "Period ID",
            "Period Name",
            "Start Date",
            "End Date",
            "Status",
            "Is Locked",
            "Business Profile ID",
            "Business Profile Name",
        ]
        ws.append(headers)

        periods = self.repository.get_all(db, skip=0, limit=2000, business_id=business_id)
        biz_repo = BusinessRepository()
        businesses = {b.id: b.name_en for b in biz_repo.get_all(db, skip=0, limit=1000)}

        for p in periods:
            ws.append([
                str(p.id),
                p.name or "",
                p.start_date.strftime("%Y-%m-%d") if p.start_date else "",
                p.end_date.strftime("%Y-%m-%d") if p.end_date else "",
                p.status.value if p.status else "",
                "Yes" if p.is_locked else "No",
                str(p.business_id) if p.business_id else "",
                businesses.get(p.business_id, "") if p.business_id else "Global",
            ])

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        output = BytesIO()
        wb.save(output)
        return output.getvalue()


class PayrollRecordService:
    def __init__(
        self,
        repository: PayrollRecordRepository | None = None,
        period_repository: PayrollPeriodRepository | None = None,
        employee_repository: EmployeeRepository | None = None,
        salary_repository: EmployeeSalaryRepository | None = None,
        attendance_repository: AttendanceRepository | None = None,
        leave_repository: LeaveApplicationRepository | None = None,
        holiday_repository: HolidayRepository | None = None,
        settings_repository: PayrollSettingsRepository | None = None,
    ):
        self.repository = repository or PayrollRecordRepository()
        self.period_repository = period_repository or PayrollPeriodRepository()
        self.employee_repository = employee_repository or EmployeeRepository()
        self.salary_repository = salary_repository or EmployeeSalaryRepository()
        self.attendance_repository = attendance_repository or AttendanceRepository()
        self.leave_repository = leave_repository or LeaveApplicationRepository()
        self.holiday_repository = holiday_repository or HolidayRepository()
        self.settings_repository = settings_repository or PayrollSettingsRepository()

    def _compute_salary_totals(
        self,
        basic_salary: float,
        house_rent: float,
        transport_allowance: float,
        medical_allowance: float,
        food_allowance: float,
        other_allowance: float,
        overtime_pay: float,
        bonus: float,
        tax: float,
        provident_fund: float,
        unpaid_leave_deduction: float,
        loan_installment: float,
        other_deduction: float,
    ) -> tuple[float, float, float]:
        gross = (
            float(basic_salary)
            + float(house_rent)
            + float(transport_allowance)
            + float(medical_allowance)
            + float(food_allowance)
            + float(other_allowance)
            + float(overtime_pay)
            + float(bonus)
        )
        deductions = (
            float(tax)
            + float(provident_fund)
            + float(unpaid_leave_deduction)
            + float(loan_installment)
            + float(other_deduction)
        )
        net = gross - deductions
        return round(gross, 2), round(deductions, 2), round(net, 2)

    def get_records(
        self,
        db: Session,
        skip: int = 0,
        limit: int = 100,
        business_id: int | None = None,
        period_id: uuid.UUID | None = None,
        employee_id: uuid.UUID | None = None,
        is_paid: bool | None = None,
    ) -> list[PayrollRecord]:
        return self.repository.get_all(
            db,
            skip=skip,
            limit=limit,
            business_id=business_id,
            period_id=period_id,
            employee_id=employee_id,
            is_paid=is_paid,
        )

    def get_record(self, db: Session, record_uuid: uuid.UUID) -> PayrollRecord:
        record = self.repository.get_by_id(db, record_uuid)
        if not record:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Payroll record (payslip) not found",
            )
        return record

    def create_record(self, db: Session, data: PayrollRecordCreate) -> PayrollRecord:
        period = self.period_repository.get_by_id(db, data.period_id)
        if not period:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Payroll period with id '{data.period_id}' not found",
            )

        if period.is_locked:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Cannot create a payslip in a locked or paid payroll period",
            )

        employee = self.employee_repository.get_by_id(db, data.employee_id)
        if not employee:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Employee with id '{data.employee_id}' not found",
            )

        existing = self.repository.get_by_period_employee(
            db, period_id=data.period_id, employee_id=data.employee_id
        )
        if existing:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Payslip record for this employee and period already exists",
            )

        gross, deductions, net = self._compute_salary_totals(
            basic_salary=data.basic_salary,
            house_rent=data.house_rent,
            transport_allowance=data.transport_allowance,
            medical_allowance=data.medical_allowance,
            food_allowance=data.food_allowance,
            other_allowance=data.other_allowance,
            overtime_pay=data.overtime_pay,
            bonus=data.bonus,
            tax=data.tax,
            provident_fund=data.provident_fund,
            unpaid_leave_deduction=data.unpaid_leave_deduction,
            loan_installment=data.loan_installment,
            other_deduction=data.other_deduction,
        )

        return self.repository.create(
            db,
            data=data,
            gross_salary=gross,
            total_deduction=deductions,
            net_salary=net,
        )

    def update_record(
        self,
        db: Session,
        record_uuid: uuid.UUID,
        data: PayrollRecordUpdate,
    ) -> PayrollRecord:
        record = self.get_record(db, record_uuid)

        if record.period and record.period.is_locked:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Cannot modify a payslip in a locked or paid payroll period",
            )

        basic_salary = (
            data.basic_salary if data.basic_salary is not None else record.basic_salary
        )
        house_rent = (
            data.house_rent if data.house_rent is not None else record.house_rent
        )
        transport_allowance = (
            data.transport_allowance
            if data.transport_allowance is not None
            else record.transport_allowance
        )
        medical_allowance = (
            data.medical_allowance
            if data.medical_allowance is not None
            else record.medical_allowance
        )
        food_allowance = (
            data.food_allowance
            if data.food_allowance is not None
            else record.food_allowance
        )
        other_allowance = (
            data.other_allowance
            if data.other_allowance is not None
            else record.other_allowance
        )
        overtime_pay = (
            data.overtime_pay if data.overtime_pay is not None else record.overtime_pay
        )
        bonus = data.bonus if data.bonus is not None else record.bonus

        tax = data.tax if data.tax is not None else record.tax
        provident_fund = (
            data.provident_fund
            if data.provident_fund is not None
            else record.provident_fund
        )
        unpaid_leave_deduction = (
            data.unpaid_leave_deduction
            if data.unpaid_leave_deduction is not None
            else record.unpaid_leave_deduction
        )
        loan_installment = (
            data.loan_installment
            if data.loan_installment is not None
            else record.loan_installment
        )
        other_deduction = (
            data.other_deduction
            if data.other_deduction is not None
            else record.other_deduction
        )

        gross, deductions, net = self._compute_salary_totals(
            basic_salary=basic_salary,
            house_rent=house_rent,
            transport_allowance=transport_allowance,
            medical_allowance=medical_allowance,
            food_allowance=food_allowance,
            other_allowance=other_allowance,
            overtime_pay=overtime_pay,
            bonus=bonus,
            tax=tax,
            provident_fund=provident_fund,
            unpaid_leave_deduction=unpaid_leave_deduction,
            loan_installment=loan_installment,
            other_deduction=other_deduction,
        )

        return self.repository.update(
            db,
            record=record,
            data=data,
            gross_salary=gross,
            total_deduction=deductions,
            net_salary=net,
        )

    def delete_record(self, db: Session, record_uuid: uuid.UUID) -> None:
        record = self.get_record(db, record_uuid)
        if record.period and record.period.is_locked:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Cannot delete a payslip in a locked or paid payroll period",
            )
        self.repository.delete(db, record)

    def generate_export_excel(self, db: Session, business_id: int | None = None) -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Payslips"

        headers = [
            "Payslip ID",
            "Period",
            "Employee ID",
            "Employee Name",
            "Basic Salary",
            "Gross Salary",
            "Total Deduction",
            "Net Salary",
            "Payment Method",
            "Is Paid",
            "Business Profile ID",
            "Business Profile Name",
        ]
        ws.append(headers)

        records = self.repository.get_all(db, skip=0, limit=2000, business_id=business_id)
        employees = {e.id: e for e in self.employee_repository.get_all(db, limit=2000)}
        periods = {p.id: p.name for p in self.period_repository.get_all(db, limit=1000)}
        biz_repo = BusinessRepository()
        businesses = {b.id: b.name_en for b in biz_repo.get_all(db, skip=0, limit=1000)}

        for r in records:
            emp = employees.get(r.employee_id)
            emp_code = emp.employee_id if emp else ""
            emp_name = emp.full_name if emp else ""

            ws.append([
                str(r.id),
                periods.get(r.period_id, ""),
                emp_code,
                emp_name,
                float(r.basic_salary or 0.0),
                float(r.gross_salary or 0.0),
                float(r.total_deduction or 0.0),
                float(r.net_salary or 0.0),
                r.payment_method.value if r.payment_method else "",
                "Yes" if r.is_paid else "No",
                str(r.business_id) if r.business_id else "",
                businesses.get(r.business_id, "") if r.business_id else "Global",
            ])

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        output = BytesIO()
        wb.save(output)
        return output.getvalue()

    def get_payroll_summary_report(
        self, db: Session, business_id: int, period_id: uuid.UUID
    ) -> dict:
        # Added for MCP read-only report aggregation (totals only, no individual employee records)
        period = self.period_repository.get_by_id(db, period_id)
        if not period:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Payroll period not found",
            )
        if period.business_id != business_id:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Payroll period does not belong to specified business profile",
            )

        records = self.repository.get_all(
            db, business_id=business_id, period_id=period_id, limit=5000
        )

        total_gross = sum(float(r.gross_salary or 0) for r in records)
        total_net = sum(float(r.net_salary or 0) for r in records)
        total_deductions = sum(float(r.total_deduction or 0) for r in records)
        total_basic = sum(float(r.basic_salary or 0) for r in records)
        total_overtime_pay = sum(float(r.overtime_pay or 0) for r in records)
        total_bonus = sum(float(r.bonus or 0) for r in records)
        paid_count = sum(1 for r in records if r.is_paid)

        return {
            "business_id": business_id,
            "period_id": str(period_id),
            "period_name": period.name,
            "period_status": str(period.status.value if hasattr(period.status, "value") else period.status),
            "headcount": len(records),
            "total_basic_salary": round(total_basic, 2),
            "total_gross_salary": round(total_gross, 2),
            "total_net_salary": round(total_net, 2),
            "total_deductions": round(total_deductions, 2),
            "total_overtime_pay": round(total_overtime_pay, 2),
            "total_bonus": round(total_bonus, 2),
            "paid_count": paid_count,
            "unpaid_count": len(records) - paid_count,
        }

    def generate_period_payroll(
        self, db: Session, period_uuid: uuid.UUID
    ) -> list[PayrollRecord]:
        period = self.period_repository.get_by_id(db, period_uuid)
        if not period:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Payroll period not found",
            )

        if period.is_locked:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Cannot generate payroll for a locked or paid payroll period",
            )

        # Get business payroll settings
        settings = self.settings_repository.get_by_business_id(
            db, business_id=period.business_id
        )
        if not settings:
            settings = self.settings_repository.create(
                db, business_id=period.business_id
            )

        # Total working days in period
        working_days = (period.end_date - period.start_date).days + 1

        # Holidays in period
        holiday_days = 0
        if settings.include_holidays:
            holidays = self.holiday_repository.get_all(
                db, business_id=period.business_id, limit=500
            )
            for h in holidays:
                overlap_start = max(period.start_date, h.start_date)
                overlap_end = min(period.end_date, h.end_date)
                if overlap_start <= overlap_end:
                    holiday_days += (overlap_end - overlap_start).days + 1

        employees = self.employee_repository.get_all(
            db, business_id=period.business_id, limit=1000
        )
        active_employees = [e for e in employees if getattr(e, "is_active", True)]

        generated_records: list[PayrollRecord] = []

        for emp in active_employees:
            salary = self.salary_repository.get_by_employee_id(
                db, employee_id=emp.id, business_id=period.business_id
            )

            basic_salary = float(getattr(salary, "basic_salary", 0.0) or 0.0)
            house_rent = float(getattr(salary, "house_rent", 0.0) or 0.0)
            medical_allowance = float(getattr(salary, "medical_allowance", 0.0) or 0.0)
            transport_allowance = float(
                getattr(salary, "transport_allowance", 0.0) or 0.0
            )
            food_allowance = float(getattr(salary, "food_allowance", 0.0) or 0.0)
            other_allowance = float(getattr(salary, "other_allowance", 0.0) or 0.0)

            tax = float(getattr(salary, "tax", 0.0) or 0.0)
            provident_fund = float(getattr(salary, "provident_fund", 0.0) or 0.0)
            other_deduction = float(getattr(salary, "other_deduction", 0.0) or 0.0)

            # Attendance processing
            present_days = 0
            att_holiday_days = 0
            overtime_hours = 0.0
            if settings.include_attendance:
                att_records = self.attendance_repository.get_all(
                    db,
                    business_id=period.business_id,
                    employee_id=emp.id,
                    start_date=period.start_date,
                    end_date=period.end_date,
                    limit=500,
                )
                for att in att_records:
                    st = getattr(att.status, "value", att.status)
                    if st in ("present", "late"):
                        present_days += 1
                    elif st == "half_day":
                        present_days += 0.5
                    elif st in ("holiday", "weekend"):
                        att_holiday_days += 1
                    
                    if settings.include_overtime:
                        overtime_hours += float(att.overtime_hours or 0.0)
            else:
                present_days = working_days

            # Leave processing
            leave_days = 0
            unpaid_leave_days = 0
            if settings.include_leave:
                leave_apps = self.leave_repository.get_all(
                    db,
                    business_id=period.business_id,
                    employee_id=emp.id,
                    status=LeaveStatusEnum.APPROVED,
                    limit=500,
                )
                for la in leave_apps:
                    ov_start = max(period.start_date, la.start_date)
                    ov_end = min(period.end_date, la.end_date)
                    if ov_start <= ov_end:
                        days = (ov_end - ov_start).days + 1
                        leave_days += days
                        if la.leave_type and not la.leave_type.is_paid:
                            unpaid_leave_days += days

            # Combine holiday days from holiday calendar and attendance records
            emp_holiday_days = max(holiday_days, att_holiday_days) if settings.include_holidays else 0

            # Absent days calculation
            if settings.include_attendance:
                absent_days = max(
                    0, int(working_days - (present_days + leave_days + emp_holiday_days))
                )
            else:
                absent_days = 0

            # Daily and hourly rates
            daily_rate = basic_salary / working_days if working_days > 0 else 0.0
            std_hrs = float(settings.standard_hours_per_day or 9.0)
            hourly_rate = daily_rate / std_hrs if std_hrs > 0 else 0.0

            # Overtime pay
            overtime_pay = 0.0
            if settings.include_overtime and overtime_hours > 0:
                overtime_pay = round(overtime_hours * hourly_rate * 1.5, 2)

            # Deductions
            unpaid_leave_ded = 0.0
            if settings.include_leave and unpaid_leave_days > 0:
                unpaid_leave_ded += unpaid_leave_days * daily_rate

            absent_ded = 0.0
            if settings.include_attendance and settings.deduct_absent_days and absent_days > 0:
                absent_ded += absent_days * daily_rate

            total_unpaid_leave_deduction = round(unpaid_leave_ded + absent_ded, 2)

            gross, deductions, net = self._compute_salary_totals(
                basic_salary=basic_salary,
                house_rent=house_rent,
                transport_allowance=transport_allowance,
                medical_allowance=medical_allowance,
                food_allowance=food_allowance,
                other_allowance=other_allowance,
                overtime_pay=overtime_pay,
                bonus=0.0,
                tax=tax,
                provident_fund=provident_fund,
                unpaid_leave_deduction=total_unpaid_leave_deduction,
                loan_installment=0.0,
                other_deduction=other_deduction,
            )

            existing = self.repository.get_by_period_employee(
                db, period_id=period.id, employee_id=emp.id
            )

            if existing:
                existing.working_days = working_days
                existing.present_days = int(present_days)
                existing.absent_days = absent_days
                existing.leave_days = leave_days
                existing.holiday_days = emp_holiday_days
                existing.overtime_hours = overtime_hours
                existing.basic_salary = basic_salary
                existing.house_rent = house_rent
                existing.medical_allowance = medical_allowance
                existing.transport_allowance = transport_allowance
                existing.food_allowance = food_allowance
                existing.other_allowance = other_allowance
                existing.overtime_pay = overtime_pay
                existing.tax = tax
                existing.provident_fund = provident_fund
                existing.unpaid_leave_deduction = total_unpaid_leave_deduction
                existing.other_deduction = other_deduction
                existing.gross_salary = gross
                existing.total_deduction = deductions
                existing.net_salary = net
                db.commit()
                db.refresh(existing)
                generated_records.append(existing)
            else:
                record_data = PayrollRecordCreate(
                    business_id=period.business_id,
                    period_id=period.id,
                    employee_id=emp.id,
                    working_days=working_days,
                    present_days=int(present_days),
                    absent_days=absent_days,
                    leave_days=leave_days,
                    holiday_days=emp_holiday_days,
                    overtime_hours=overtime_hours,
                    basic_salary=basic_salary,
                    house_rent=house_rent,
                    transport_allowance=transport_allowance,
                    medical_allowance=medical_allowance,
                    food_allowance=food_allowance,
                    other_allowance=other_allowance,
                    overtime_pay=overtime_pay,
                    bonus=0.0,
                    tax=tax,
                    provident_fund=provident_fund,
                    unpaid_leave_deduction=total_unpaid_leave_deduction,
                    loan_installment=0.0,
                    other_deduction=other_deduction,
                    payment_method=PaymentMethodEnum.BANK_TRANSFER,
                    is_paid=False,
                    note=f"Auto-generated salary for {period.name}",
                )
                new_rec = self.repository.create(
                    db,
                    data=record_data,
                    gross_salary=gross,
                    total_deduction=deductions,
                    net_salary=net,
                )
                generated_records.append(new_rec)

        return generated_records


class PayrollSettingsService:
    def __init__(self, repository: PayrollSettingsRepository | None = None):
        self.repository = repository or PayrollSettingsRepository()

    def get_settings(self, db: Session, business_id: int) -> PayrollSettings:
        settings_obj = self.repository.get_by_business_id(db, business_id)
        if not settings_obj:
            settings_obj = self.repository.create(db, business_id=business_id)
        return settings_obj

    def update_settings(
        self, db: Session, business_id: int, data: PayrollSettingsUpdate
    ) -> PayrollSettings:
        settings_obj = self.repository.get_by_business_id(db, business_id)
        if not settings_obj:
            settings_obj = self.repository.create(db, business_id=business_id, data=data)
        else:
            settings_obj = self.repository.update(db, settings_obj, data)
        return settings_obj

    def generate_export_excel(self, db: Session, business_id: int | None = None) -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Payroll Settings"

        headers = [
            "Settings ID",
            "Business Profile ID",
            "Business Profile Name",
            "Standard Hours Per Day",
            "Include Attendance",
            "Include Overtime",
            "Include Leave",
            "Include Holidays",
            "Deduct Absent Days",
        ]
        ws.append(headers)

        biz_repo = BusinessRepository()
        businesses = biz_repo.get_all(db, skip=0, limit=1000)
        target_businesses = [b for b in businesses if business_id is None or b.id == business_id]

        for b in target_businesses:
            st = self.get_settings(db, business_id=b.id)
            ws.append([
                str(st.id),
                str(b.id),
                b.name_en or "",
                float(st.standard_hours_per_day or 9.0),
                "Yes" if st.include_attendance else "No",
                "Yes" if st.include_overtime else "No",
                "Yes" if st.include_leave else "No",
                "Yes" if st.include_holidays else "No",
                "Yes" if st.deduct_absent_days else "No",
            ])

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        output = BytesIO()
        wb.save(output)
        return output.getvalue()
