import uuid
from typing import Any
from io import BytesIO
import openpyxl
import re

from fastapi import HTTPException, status
from sqlalchemy.orm import Session

from app.core.tenancy.models import BusinessProfile
from app.core.tenancy.repository import BusinessRepository
from app.core.tenancy.scoping import verify_record_ownership
from app.modules.hr_payroll.employees.models import Employee
from app.modules.hr_payroll.organization.models import Branch, Department, JobTitle
from app.modules.hr_payroll.organization.repository import (
    BranchRepository,
    DepartmentRepository,
    JobTitleRepository,
)
from app.modules.hr_payroll.organization.schemas import (
    BranchCreate,
    BranchOut,
    BranchStructureNode,
    BranchUpdate,
    DepartmentCreate,
    DepartmentTreeNode,
    DepartmentUpdate,
    JobTitleCreate,
    JobTitleUpdate,
    OrganogramEmployeeNode,
    OrganogramStructureTree,
)


class DepartmentService:

    def __init__(self, repository: DepartmentRepository | None = None):
        self.repository = repository or DepartmentRepository()

    def get_departments(
        self,
        db: Session,
        skip: int = 0,
        limit: int = 100,
        business_id: int | None = None,
        search: str | None = None,
        is_active: bool | None = None,
        parent_id: uuid.UUID | None = None,
    ) -> list[Department]:
        return self.repository.get_all(
            db,
            skip=skip,
            limit=limit,
            business_id=business_id,
            search=search,
            is_active=is_active,
            parent_id=parent_id,
        )

    def get_department_tree(
        self, db: Session, business_id: int | None = None
    ) -> list[DepartmentTreeNode]:
        departments = self.repository.get_all(db, business_id=business_id, limit=500)
        nodes: dict[uuid.UUID, DepartmentTreeNode] = {}
        for d in departments:
            nodes[d.id] = DepartmentTreeNode(
                id=d.id,
                name=d.name,
                code=d.code,
                slug=d.slug,
                description=d.description,
                is_active=d.is_active,
                head_employee_id=d.head_employee_id,
                head_employee_name=d.head_employee_name,
                employee_count=d.employee_count,
                children=[],
            )

        tree: list[DepartmentTreeNode] = []
        for d in departments:
            node = nodes[d.id]
            if d.parent_department_id and d.parent_department_id in nodes:
                nodes[d.parent_department_id].children.append(node)
            else:
                tree.append(node)
        return tree

    def get_department(
        self, db: Session, department_id: uuid.UUID, current_user: Any | None = None
    ) -> Department:
        department = self.repository.get_by_id(db, department_id)
        if current_user is not None:
            return verify_record_ownership(department, current_user, detail="Department not found")
        if not department:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Department not found",
            )
        return department

    def create_department(self, db: Session, data: DepartmentCreate) -> Department:
        if self.repository.get_by_slug(db, data.slug, data.business_id):
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Department with slug '{data.slug}' already exists",
            )
        if self.repository.get_by_name(db, data.name, data.business_id):
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Department with name '{data.name}' already exists",
            )

        return self.repository.create(db, data)

    def update_department(
        self,
        db: Session,
        department_id: uuid.UUID,
        data: DepartmentUpdate,
        current_user: Any | None = None,
    ) -> Department:
        department = self.get_department(db, department_id, current_user=current_user)

        target_business_id = (
            data.business_id
            if data.business_id is not None
            else department.business_id
        )

        if data.slug is not None and data.slug != department.slug:
            existing = self.repository.get_by_slug(
                db, data.slug, target_business_id
            )
            if existing and existing.id != department_id:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"Department with slug '{data.slug}' already exists",
                )

        if data.name is not None and data.name != department.name:
            existing_name = self.repository.get_by_name(
                db, data.name, target_business_id
            )
            if existing_name and existing_name.id != department_id:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"Department with name '{data.name}' already exists",
                )

        if data.parent_department_id is not None:
            if data.parent_department_id == department_id:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="A department cannot be its own parent department.",
                )
            descendant_ids = self.repository.get_descendant_ids(db, department_id)
            if data.parent_department_id in descendant_ids:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="Cannot assign a sub-department as parent (cyclic hierarchy dependency).",
                )

        return self.repository.update(db, department, data)

    def delete_department(
        self, db: Session, department_id: uuid.UUID, current_user: Any | None = None
    ) -> None:
        department = self.get_department(db, department_id, current_user=current_user)
        self.repository.delete(db, department)

    def generate_export_excel(self, db: Session, business_id: int | None = None) -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Departments"

        headers = [
            "Department ID",
            "Department Name",
            "Slug",
            "Description",
            "Business Profile ID",
            "Business Profile Name",
            "Multiple Heads Allowed",
            "Status",
        ]
        ws.append(headers)

        departments = self.repository.get_all(db, skip=0, limit=2000, business_id=business_id)
        biz_repo = BusinessRepository()
        businesses = {b.id: b.name_en for b in biz_repo.get_all(db, skip=0, limit=1000)}

        for d in departments:
            ws.append([
                str(d.id),
                d.name or "",
                d.slug or "",
                d.description or "",
                str(d.business_id) if d.business_id else "",
                businesses.get(d.business_id, "") if d.business_id else "Global",
                "Yes" if d.multiple_heads_allowed else "No",
                "Active" if d.is_active else "Inactive",
            ])

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        output = BytesIO()
        wb.save(output)
        return output.getvalue()

    def generate_excel_template(self, db: Session, business_id: int) -> bytes:
        if not business_id or business_id <= 0:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Business profile ID is mandatory.",
            )
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Department Template"

        headers = [
            "Department Name",
            "Slug",
            "Description",
            "Multiple Heads Allowed (yes/no)",
        ]
        ws.append(headers)

        ws.append([
            "Engineering",
            "engineering",
            "Software engineering and technology department",
            "no",
        ])
        ws.append([
            "Human Resources",
            "human-resources",
            "HR and talent management department",
            "no",
        ])

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

        output = BytesIO()
        wb.save(output)
        return output.getvalue()

    def import_departments_excel(
        self, db: Session, business_id: int, file_bytes: bytes
    ) -> dict[str, int | list[str]]:
        if not business_id or business_id <= 0:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Business profile ID is mandatory.",
            )
        try:
            wb = openpyxl.load_workbook(filename=BytesIO(file_bytes), data_only=True)
            ws = wb.active
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Invalid Excel file format: {str(e)}",
            )

        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Excel file is empty.",
            )

        success_count = 0
        error_messages: list[str] = []

        header = [str(cell or "").strip().lower() for cell in rows[0]] if rows else []
        has_header = any(h in header for h in ["department name", "name", "slug", "description"])
        start_idx = 1 if has_header else 0

        col_map = {}
        if has_header:
            for idx, h in enumerate(header):
                if h in ["department name", "name"]:
                    col_map["name"] = idx
                elif h in ["slug"]:
                    col_map["slug"] = idx
                elif h in ["description"]:
                    col_map["description"] = idx
                elif "multiple heads" in h or h in ["multiple heads allowed", "multiple_heads_allowed"]:
                    col_map["multi_head"] = idx

        for row_idx, row in enumerate(rows[start_idx:], start=start_idx + 1):
            if not row or not any(row):
                continue

            def get_val(key, default_idx):
                idx = col_map.get(key, default_idx)
                if idx is not None and idx < len(row) and row[idx] is not None:
                    return row[idx]
                return None

            name_raw = str(get_val("name", 0) or "").strip()
            slug_raw = str(get_val("slug", 1) or "").strip().lower() if get_val("slug", 1) else ""
            desc_raw = str(get_val("description", 2) or "").strip() if get_val("description", 2) else None
            multi_head_raw = str(get_val("multi_head", 3) or "no").strip().lower()

            if not name_raw:
                error_messages.append(f"Row {row_idx}: Missing Department Name.")
                continue

            if not slug_raw:
                slug_raw = re.sub(r"[^\w\s-]", "", name_raw).strip().lower().replace(" ", "-")

            # Check duplicate name or slug
            existing_slug = self.repository.get_by_slug(db, slug_raw, business_id)
            existing_name = self.repository.get_by_name(db, name_raw, business_id)

            if existing_slug or existing_name:
                warning_reason = []
                if existing_name:
                    warning_reason.append(f"Department name '{name_raw}' already exists")
                if existing_slug:
                    warning_reason.append(f"Slug '{slug_raw}' already exists")
                error_messages.append(f"Row {row_idx}: Skipped - {', '.join(warning_reason)}.")
                continue

            multi_head = multi_head_raw in ("yes", "true", "1")

            create_data = DepartmentCreate(
                business_id=business_id,
                name=name_raw,
                slug=slug_raw,
                description=desc_raw,
                is_active=True,
                multiple_heads_allowed=multi_head,
            )

            try:
                self.create_department(db, create_data)
                success_count += 1
            except HTTPException as hexp:
                error_messages.append(f"Row {row_idx}: {hexp.detail}")
            except Exception as ex:
                error_messages.append(f"Row {row_idx}: Failed to create department - {str(ex)}")

        return {
            "imported_count": success_count,
            "errors": error_messages,
        }


class JobTitleService:

    def __init__(
        self,
        repository: JobTitleRepository | None = None,
        department_repository: DepartmentRepository | None = None,
    ):
        self.repository = repository or JobTitleRepository()
        self.department_repository = (
            department_repository or DepartmentRepository()
        )

    def get_job_titles(
        self,
        db: Session,
        skip: int = 0,
        limit: int = 100,
        business_id: int | None = None,
        department_id: uuid.UUID | None = None,
        search: str | None = None,
        is_active: bool | None = None,
        level_rank: int | None = None,
    ) -> list[JobTitle]:
        return self.repository.get_all(
            db,
            skip=skip,
            limit=limit,
            business_id=business_id,
            department_id=department_id,
            search=search,
            is_active=is_active,
            level_rank=level_rank,
        )

    def get_job_title(
        self, db: Session, job_title_id: uuid.UUID, current_user: Any | None = None
    ) -> JobTitle:
        job_title = self.repository.get_by_id(db, job_title_id)
        if current_user is not None:
            return verify_record_ownership(job_title, current_user, detail="Job title not found")
        if not job_title:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Job title not found",
            )
        return job_title

    def create_job_title(self, db: Session, data: JobTitleCreate) -> JobTitle:
        if data.department_id is not None:
            dept = self.department_repository.get_by_id(db, data.department_id)
            if not dept:
                raise HTTPException(
                    status_code=status.HTTP_404_NOT_FOUND,
                    detail=f"Department with id '{data.department_id}' not found",
                )
            existing = self.repository.get_by_name(
                db, data.name, department_id=data.department_id
            )
            if existing:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=(
                        f"Job title with name '{data.name}' already exists in"
                        " this department"
                    ),
                )

        return self.repository.create(db, data)

    def update_job_title(
        self,
        db: Session,
        job_title_id: uuid.UUID,
        data: JobTitleUpdate,
        current_user: Any | None = None,
    ) -> JobTitle:
        job_title = self.get_job_title(db, job_title_id, current_user=current_user)

        target_department_id = (
            data.department_id
            if data.department_id is not None
            else job_title.department_id
        )

        if target_department_id is not None:
            dept = self.department_repository.get_by_id(db, target_department_id)
            if not dept:
                raise HTTPException(
                    status_code=status.HTTP_404_NOT_FOUND,
                    detail=f"Department with id '{target_department_id}' not found",
                )

            target_name = data.name if data.name is not None else job_title.name
            existing = self.repository.get_by_name(
                db, target_name, department_id=target_department_id
            )
            if existing and existing.id != job_title_id:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=(
                        f"Job title with name '{target_name}' already exists in"
                        " this department"
                    ),
                )

        return self.repository.update(db, job_title, data)

    def delete_job_title(
        self, db: Session, job_title_id: uuid.UUID, current_user: Any | None = None
    ) -> None:
        job_title = self.get_job_title(db, job_title_id, current_user=current_user)
        self.repository.delete(db, job_title)

    def generate_export_excel(
        self,
        db: Session,
        business_id: int | None = None,
        department_id: uuid.UUID | None = None,
    ) -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Job Titles"

        headers = [
            "Job Title ID",
            "Job Title Name",
            "Short Name",
            "Department",
            "Business Profile ID",
            "Business Profile Name",
            "Description",
            "Status",
        ]
        ws.append(headers)

        job_titles = self.repository.get_all(
            db, skip=0, limit=2000, business_id=business_id, department_id=department_id
        )
        departments = {d.id: d.name for d in self.department_repository.get_all(db, limit=1000)}
        biz_repo = BusinessRepository()
        businesses = {b.id: b.name_en for b in biz_repo.get_all(db, skip=0, limit=1000)}

        for j in job_titles:
            ws.append([
                str(j.id),
                j.name or "",
                j.short_name or "",
                departments.get(j.department_id, "") if j.department_id else "Unassigned",
                str(j.business_id) if j.business_id else "",
                businesses.get(j.business_id, "") if j.business_id else "Global",
                j.description or "",
                "Active" if j.is_active else "Inactive",
            ])

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        output = BytesIO()
        wb.save(output)
        return output.getvalue()

    def generate_excel_template(self, db: Session, business_id: int) -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Job Title Template"

        headers = [
            "Job Title Name",
            "Short Name",
            "Department Name",
            "Description",
        ]
        ws.append(headers)

        departments = self.department_repository.get_all(db, business_id=business_id, limit=10)
        sample_dept = departments[0].name if departments else "Engineering"

        ws.append([
            "Senior Software Engineer",
            "Sr. SE",
            sample_dept,
            "Lead software engineer role",
        ])
        ws.append([
            "Product Manager",
            "PM",
            sample_dept,
            "Product management role",
        ])

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

        output = BytesIO()
        wb.save(output)
        return output.getvalue()

    def import_job_titles_excel(
        self, db: Session, business_id: int, file_bytes: bytes
    ) -> dict[str, int | list[str]]:
        try:
            wb = openpyxl.load_workbook(filename=BytesIO(file_bytes), data_only=True)
            ws = wb.active
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Invalid Excel file format: {str(e)}",
            )

        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Excel file is empty.",
            )

        departments = self.department_repository.get_all(db, business_id=business_id, limit=500)
        dept_map = {d.name.strip().lower(): d for d in departments if d.name}

        success_count = 0
        error_messages: list[str] = []

        header = [str(cell or "").strip().lower() for cell in rows[0]]
        start_idx = 1 if "job title name" in header or "name" in header or "short name" in header else 0

        for row_idx, row in enumerate(rows[start_idx:], start=start_idx + 1):
            if not row or not any(row):
                continue

            name_raw = str(row[0] or "").strip()
            short_name_raw = str(row[1] or "").strip() if len(row) > 1 and row[1] else None
            dept_raw = str(row[2] or "").strip() if len(row) > 2 and row[2] else None
            desc_raw = str(row[3] or "").strip() if len(row) > 3 and row[3] else None

            if not name_raw:
                error_messages.append(f"Row {row_idx}: Missing Job Title Name.")
                continue

            dept_obj = dept_map.get(dept_raw.lower()) if dept_raw else None
            dept_id = dept_obj.id if dept_obj else None

            if dept_raw and not dept_obj:
                error_messages.append(f"Row {row_idx}: Department '{dept_raw}' not found.")
                continue

            if dept_id:
                existing = self.repository.get_by_name(db, name_raw, department_id=dept_id)
                if existing:
                    error_messages.append(
                        f"Row {row_idx}: Skipped - Job title '{name_raw}' already exists in department '{dept_raw}'."
                    )
                    continue

            create_data = JobTitleCreate(
                business_id=business_id,
                department_id=dept_id,
                name=name_raw,
                short_name=short_name_raw,
                description=desc_raw,
                is_active=True,
            )

            try:
                self.create_job_title(db, create_data)
                success_count += 1
            except HTTPException as hexp:
                error_messages.append(f"Row {row_idx}: {hexp.detail}")
            except Exception as ex:
                error_messages.append(f"Row {row_idx}: Failed to create job title - {str(ex)}")

        return {
            "imported_count": success_count,
            "errors": error_messages,
        }


class BranchService:

    def __init__(self, repository: BranchRepository | None = None):
        self.repository = repository or BranchRepository()
        self.business_repo = BusinessRepository()

    def get_branches(
        self,
        db: Session,
        skip: int = 0,
        limit: int = 100,
        business_id: int | None = None,
        search: str | None = None,
        is_active: bool | None = None,
        branch_type: str | None = None,
    ) -> list[Branch]:
        return self.repository.get_all(
            db,
            skip=skip,
            limit=limit,
            business_id=business_id,
            search=search,
            is_active=is_active,
            branch_type=branch_type,
        )

    def count_branches(
        self,
        db: Session,
        business_id: int | None = None,
        search: str | None = None,
        is_active: bool | None = None,
    ) -> int:
        return self.repository.count(
            db, business_id=business_id, search=search, is_active=is_active
        )

    def get_branch(
        self, db: Session, branch_id: uuid.UUID, current_user: Any | None = None
    ) -> Branch:
        branch = self.repository.get_by_id(db, branch_id)
        if current_user is not None:
            return verify_record_ownership(branch, current_user, detail="Branch not found")
        if not branch:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Branch not found",
            )
        return branch

    def create_branch(self, db: Session, data: BranchCreate) -> Branch:
        if self.repository.get_by_code(db, data.code, data.business_id):
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Branch with code '{data.code}' already exists for this business",
            )
        if self.repository.get_by_name(db, data.name, data.business_id):
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Branch with name '{data.name}' already exists for this business",
            )
        if data.is_head_office:
            existing_branches = self.repository.get_all(
                db, business_id=data.business_id, limit=500
            )
            for eb in existing_branches:
                if eb.is_head_office:
                    eb.is_head_office = False
            db.commit()

        return self.repository.create(db, data)

    def update_branch(
        self,
        db: Session,
        branch_id: uuid.UUID,
        data: BranchUpdate,
        current_user: Any | None = None,
    ) -> Branch:
        branch = self.get_branch(db, branch_id, current_user=current_user)
        target_business_id = (
            data.business_id if data.business_id is not None else branch.business_id
        )

        if data.code and data.code != branch.code:
            existing = self.repository.get_by_code(db, data.code, target_business_id)
            if existing and existing.id != branch_id:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"Branch with code '{data.code}' already exists",
                )

        if data.name and data.name != branch.name:
            existing = self.repository.get_by_name(db, data.name, target_business_id)
            if existing and existing.id != branch_id:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"Branch with name '{data.name}' already exists",
                )

        if data.is_head_office:
            existing_branches = self.repository.get_all(
                db, business_id=target_business_id, limit=500
            )
            for eb in existing_branches:
                if eb.id != branch_id and eb.is_head_office:
                    eb.is_head_office = False
            db.commit()

        return self.repository.update(db, branch, data)

    def delete_branch(
        self, db: Session, branch_id: uuid.UUID, current_user: Any | None = None
    ) -> None:
        branch = self.get_branch(db, branch_id, current_user=current_user)
        if branch.employees and len(branch.employees) > 0:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Cannot delete branch '{branch.name}' because {len(branch.employees)} employee(s) are assigned to it. Please reassign them first.",
            )
        self.repository.delete(db, branch)

    def generate_export_excel(self, db: Session, business_id: int | None = None) -> bytes:
        branches = self.repository.get_all(db, business_id=business_id, limit=1000)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Branches"

        headers = [
            "Name",
            "Code",
            "Type",
            "Manager",
            "Headcount",
            "Phone",
            "Email",
            "City",
            "Country",
            "Is Head Office",
            "Status",
        ]
        ws.append(headers)

        for b in branches:
            ws.append([
                b.name,
                b.code,
                b.branch_type.replace("_", " ").title(),
                b.manager_name or "Unassigned",
                b.employee_count,
                b.phone or "",
                b.email or "",
                b.city or "",
                b.country or "",
                "Yes" if b.is_head_office else "No",
                "Active" if b.is_active else "Inactive",
            ])

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        output = BytesIO()
        wb.save(output)
        return output.getvalue()


class OrganogramService:

    def __init__(
        self,
        branch_repo: BranchRepository | None = None,
        dept_repo: DepartmentRepository | None = None,
    ):
        self.branch_repo = branch_repo or BranchRepository()
        self.dept_repo = dept_repo or DepartmentRepository()

    def get_reporting_organogram(
        self,
        db: Session,
        business_id: int | None = None,
        branch_id: uuid.UUID | None = None,
        department_id: uuid.UUID | None = None,
        root_employee_id: uuid.UUID | None = None,
    ) -> list[OrganogramEmployeeNode]:
        """
        Builds a recursive tree of employees representing the reporting hierarchy (Chain of Command).
        Nodes connect via direct_manager_id. Top-level nodes are employees without a manager
        (or whose manager is outside the filtered scope).
        """
        query = db.query(Employee).filter(Employee.is_active == True)
        if business_id is not None:
            query = query.filter(Employee.business_id == business_id)
        if branch_id is not None:
            query = query.filter(Employee.branch_id == branch_id)
        if department_id is not None:
            query = query.filter(Employee.department_id == department_id)

        employees = query.all()
        if not employees:
            return []

        # Map employees to OrganogramEmployeeNode
        node_map: dict[uuid.UUID, OrganogramEmployeeNode] = {}
        for emp in employees:
            job_title_name = emp.job_title.name if emp.job_title else None
            level_rank = emp.job_title.level_rank if emp.job_title else 5
            node_map[emp.id] = OrganogramEmployeeNode(
                id=emp.id,
                employee_id=emp.employee_id,
                name=emp.full_name,
                work_email=emp.work_email,
                phone=emp.phone,
                job_title=job_title_name,
                level_rank=level_rank,
                department=emp.department_name,
                department_id=emp.department_id,
                branch=emp.branch_name,
                branch_id=emp.branch_id,
                is_department_head=emp.is_department_head,
                is_active=emp.is_active,
                direct_manager_id=emp.direct_manager_id,
                subordinate_count=0,
                subordinates=[],
            )

        # Build tree
        roots: list[OrganogramEmployeeNode] = []
        for emp in employees:
            node = node_map[emp.id]
            mgr_id = emp.direct_manager_id

            if root_employee_id is not None and emp.id == root_employee_id:
                roots.append(node)
                continue

            if root_employee_id is None:
                if mgr_id and mgr_id in node_map and mgr_id != emp.id:
                    node_map[mgr_id].subordinates.append(node)
                else:
                    roots.append(node)

        # Calculate subordinate counts recursively
        def calculate_counts(curr: OrganogramEmployeeNode) -> int:
            total = len(curr.subordinates)
            for child in curr.subordinates:
                total += calculate_counts(child)
            curr.subordinate_count = total
            return total

        for r in roots:
            calculate_counts(r)

        # Sort roots by level_rank ascending (1=top leadership), then name
        roots.sort(key=lambda x: (x.level_rank, x.name))
        return roots

    def get_structure_organogram(
        self, db: Session, business_id: int
    ) -> OrganogramStructureTree:
        """
        Builds the complete multi-branch company structure:
        Company -> Branches -> Departments -> Headcount & Staff rosters.
        """
        biz = db.query(BusinessProfile).filter(BusinessProfile.id == business_id).first()
        company_name = biz.name_en if biz else f"Company #{business_id}"

        branches = self.branch_repo.get_all(db, business_id=business_id, limit=500)
        departments = self.dept_repo.get_all(db, business_id=business_id, limit=500)
        employees = (
            db.query(Employee)
            .filter(Employee.business_id == business_id, Employee.is_active == True)
            .all()
        )

        dept_nodes = [
            DepartmentTreeNode(
                id=d.id,
                name=d.name,
                code=d.code,
                slug=d.slug,
                description=d.description,
                is_active=d.is_active,
                head_employee_id=d.head_employee_id,
                head_employee_name=d.head_employee_name,
                employee_count=d.employee_count,
                children=[],
            )
            for d in departments
        ]

        branch_nodes: list[BranchStructureNode] = []
        for b in branches:
            branch_emps = [e for e in employees if e.branch_id == b.id]
            emp_nodes = [
                OrganogramEmployeeNode(
                    id=e.id,
                    employee_id=e.employee_id,
                    name=e.full_name,
                    work_email=e.work_email,
                    phone=e.phone,
                    job_title=e.job_title_name,
                    level_rank=e.job_title.level_rank if e.job_title else 5,
                    department=e.department_name,
                    department_id=e.department_id,
                    branch=b.name,
                    branch_id=b.id,
                    is_department_head=e.is_department_head,
                    is_active=e.is_active,
                    direct_manager_id=e.direct_manager_id,
                    subordinate_count=0,
                    subordinates=[],
                )
                for e in branch_emps
            ]
            branch_nodes.append(
                BranchStructureNode(
                    branch_id=b.id,
                    branch_name=b.name,
                    branch_code=b.code,
                    branch_type=b.branch_type,
                    is_head_office=b.is_head_office,
                    city=b.city,
                    manager_name=b.manager_name,
                    employee_count=len(branch_emps),
                    departments=dept_nodes,
                    employees=emp_nodes,
                )
            )

        return OrganogramStructureTree(
            business_id=business_id,
            company_name=company_name,
            total_branches=len(branches),
            total_departments=len(departments),
            total_employees=len(employees),
            branches=branch_nodes,
        )

