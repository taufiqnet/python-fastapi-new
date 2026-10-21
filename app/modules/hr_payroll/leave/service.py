import uuid
from datetime import date, datetime, timedelta, timezone
from io import BytesIO

from fastapi import HTTPException, status
import openpyxl
from sqlalchemy.orm import Session

from app.core.tenancy.repository import BusinessRepository
from app.core.tenancy.scoping import verify_record_ownership
from app.modules.hr_payroll.attendance.models import (
    Attendance,
    AttendanceSourceEnum,
    AttendanceStatusEnum,
)
from app.modules.hr_payroll.compensation.repository import EmployeeSalaryRepository
from app.modules.hr_payroll.employees.repository import EmployeeRepository
from app.modules.hr_payroll.leave.models import (
    GenderApplicabilityEnum,
    LeaveAllocation,
    LeaveApplication,
    LeaveCarryForwardLog,
    LeaveEncashment,
    LeaveEncashmentStatusEnum,
    LeaveStatusEnum,
    LeaveType,
)
from app.modules.hr_payroll.leave.repository import (
    LeaveAllocationRepository,
    LeaveApplicationRepository,
    LeaveCarryForwardLogRepository,
    LeaveEncashmentRepository,
    LeaveTypeRepository,
)
from app.modules.hr_payroll.leave.schemas import (
    LeaveAllocationCreate,
    LeaveAllocationUpdate,
    LeaveApplicationCreate,
    LeaveApplicationReview,
    LeaveApplicationRollback,
    LeaveApplicationUpdate,
    LeaveCarryForwardExecuteRequest,
    LeaveCarryForwardExecuteResponse,
    LeaveCarryForwardPreviewItem,
    LeaveCarryForwardPreviewRequest,
    LeaveCarryForwardPreviewResponse,
    LeaveEncashmentCreate,
    LeaveEncashmentDisburse,
    LeaveEncashmentReview,
    LeaveTypeCreate,
    LeaveTypeUpdate,
)
from app.modules.hr_payroll.payroll.models import (
    PayrollPeriod,
    PayrollPeriodStatusEnum,
)


class LeaveTypeService:
    def __init__(self, repository: LeaveTypeRepository | None = None):
        self.repository = repository or LeaveTypeRepository()

    def get_leave_types(
        self,
        db: Session,
        skip: int = 0,
        limit: int = 100,
        business_id: int | None = None,
    ) -> list[LeaveType]:
        return self.repository.get_all(
            db, skip=skip, limit=limit, business_id=business_id
        )

    def get_leave_type(
        self, db: Session, leave_type_uuid: uuid.UUID, current_user=None
    ) -> LeaveType:
        leave_type = self.repository.get_by_id(db, leave_type_uuid)
        verify_record_ownership(leave_type, current_user, detail="Leave type not found")
        return leave_type

    def create_leave_type(self, db: Session, data: LeaveTypeCreate) -> LeaveType:
        if self.repository.get_by_code(db, data.code, data.business_id):
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Leave type with code '{data.code}' already exists",
            )
        return self.repository.create(db, data)

    def update_leave_type(
        self,
        db: Session,
        leave_type_uuid: uuid.UUID,
        data: LeaveTypeUpdate,
        current_user=None,
    ) -> LeaveType:
        leave_type = self.get_leave_type(db, leave_type_uuid, current_user=current_user)
        target_business_id = (
            data.business_id if data.business_id is not None else leave_type.business_id
        )

        if data.code is not None and data.code != leave_type.code:
            existing = self.repository.get_by_code(db, data.code, target_business_id)
            if existing and existing.id != leave_type_uuid:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"Leave type with code '{data.code}' already exists",
                )

        return self.repository.update(db, leave_type, data)

    def delete_leave_type(
        self, db: Session, leave_type_uuid: uuid.UUID, current_user=None
    ) -> None:
        leave_type = self.get_leave_type(db, leave_type_uuid, current_user=current_user)
        self.repository.delete(db, leave_type)

    def generate_excel_template(self, db: Session, business_id: int) -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Leave Type Template"

        headers = [
            "Name",
            "Code (Short Code e.g. AL, SL)",
            "Max Days Per Year (0 for unlimited)",
            "Is Paid (yes/no)",
            "Requires Document (yes/no)",
            "Applicable Gender (all/male/female)",
            "Carry Forward (yes/no)",
            "Description",
        ]
        ws.append(headers)

        ws.append([
            "Annual Leave",
            "AL",
            20,
            "yes",
            "no",
            "all",
            "yes",
            "Paid annual holiday entitlement",
        ])
        ws.append([
            "Sick Leave",
            "SL",
            10,
            "yes",
            "yes",
            "all",
            "no",
            "Medical sick leave",
        ])

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

        output = BytesIO()
        wb.save(output)
        return output.getvalue()

    def generate_export_excel(self, db: Session, business_id: int | None = None) -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Leave Types"

        headers = [
            "Leave Type ID",
            "Name",
            "Code",
            "Description",
            "Max Days Per Year",
            "Is Paid",
            "Requires Document",
            "Applicable Gender",
            "Carry Forward",
            "Status",
            "Business Profile ID",
            "Business Profile Name",
        ]
        ws.append(headers)

        types = self.repository.get_all(db, skip=0, limit=2000, business_id=business_id)
        biz_repo = BusinessRepository()
        businesses = {b.id: b.name_en for b in biz_repo.get_all(db, skip=0, limit=1000)}

        for t in types:
            ws.append([
                str(t.id),
                t.name or "",
                t.code or "",
                t.description or "",
                t.max_days_per_year or 0,
                "Yes" if t.is_paid else "No",
                "Yes" if t.requires_document else "No",
                t.applicable_gender.value if t.applicable_gender else "",
                "Yes" if t.carry_forward else "No",
                "Active" if t.is_active else "Inactive",
                str(t.business_id) if t.business_id else "",
                businesses.get(t.business_id, "") if t.business_id else "Global",
            ])

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        output = BytesIO()
        wb.save(output)
        return output.getvalue()

    def import_leave_types_excel(
        self, db: Session, business_id: int, file_bytes: bytes
    ) -> dict[str, int | list[str]]:
        try:
            wb = openpyxl.load_workbook(filename=BytesIO(file_bytes), data_only=True)
            ws = wb.active
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Invalid Excel file format: {str(e)}",
            )

        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Excel file is empty.",
            )

        success_count = 0
        error_messages: list[str] = []

        header = [str(cell or "").strip().lower() for cell in rows[0]]
        start_idx = 1 if "name" in header or "code" in header else 0

        for row_idx, row in enumerate(rows[start_idx:], start=start_idx + 1):
            if not row or not any(row):
                continue

            name_raw = str(row[0] or "").strip()
            code_raw = str(row[1] or "").strip().upper() if len(row) > 1 and row[1] else ""
            max_days_raw = row[2] if len(row) > 2 else 0
            is_paid_raw = str(row[3] or "yes").strip().lower() if len(row) > 3 and row[3] is not None else "yes"
            req_doc_raw = str(row[4] or "no").strip().lower() if len(row) > 4 and row[4] is not None else "no"
            gender_raw = str(row[5] or "all").strip().lower() if len(row) > 5 and row[5] else "all"
            carry_fwd_raw = str(row[6] or "no").strip().lower() if len(row) > 6 and row[6] is not None else "no"
            desc_raw = str(row[7] or "").strip() if len(row) > 7 and row[7] else None

            if not name_raw:
                error_messages.append(f"Row {row_idx}: Missing Leave Type Name.")
                continue

            if not code_raw:
                error_messages.append(f"Row {row_idx}: Missing Leave Type Code.")
                continue

            # Check existing code for warning message
            existing = self.repository.get_by_code(db, code_raw, business_id)
            if existing:
                error_messages.append(f"Row {row_idx}: Skipped - Leave type code '{code_raw}' already exists.")
                continue

            # Max days parsing
            try:
                max_days = int(max_days_raw) if max_days_raw is not None and str(max_days_raw).strip() != "" else 0
            except (ValueError, TypeError):
                max_days = 0

            is_paid = is_paid_raw in ("yes", "true", "1")
            requires_document = req_doc_raw in ("yes", "true", "1")
            carry_forward = carry_fwd_raw in ("yes", "true", "1")

            try:
                gender_enum = GenderApplicabilityEnum(gender_raw)
            except ValueError:
                gender_enum = GenderApplicabilityEnum.ALL

            create_data = LeaveTypeCreate(
                business_id=business_id,
                name=name_raw,
                code=code_raw,
                description=desc_raw,
                max_days_per_year=max_days,
                is_paid=is_paid,
                requires_document=requires_document,
                applicable_gender=gender_enum,
                carry_forward=carry_forward,
                is_active=True,
            )

            try:
                self.create_leave_type(db, create_data)
                success_count += 1
            except HTTPException as hexp:
                error_messages.append(f"Row {row_idx}: {hexp.detail}")
            except Exception as ex:
                error_messages.append(f"Row {row_idx}: Failed to create leave type - {str(ex)}")

        return {
            "imported_count": success_count,
            "errors": error_messages,
        }


class LeaveAllocationService:
    def __init__(
        self,
        repository: LeaveAllocationRepository | None = None,
        leave_type_repository: LeaveTypeRepository | None = None,
        employee_repository: EmployeeRepository | None = None,
    ):
        self.repository = repository or LeaveAllocationRepository()
        self.leave_type_repository = leave_type_repository or LeaveTypeRepository()
        self.employee_repository = employee_repository or EmployeeRepository()

    def get_allocations(
        self,
        db: Session,
        skip: int = 0,
        limit: int = 100,
        business_id: int | None = None,
        employee_id: uuid.UUID | None = None,
        year: int | None = None,
    ) -> list[LeaveAllocation]:
        return self.repository.get_all(
            db,
            skip=skip,
            limit=limit,
            business_id=business_id,
            employee_id=employee_id,
            year=year,
        )

    def get_allocation(
        self, db: Session, allocation_uuid: uuid.UUID, current_user=None
    ) -> LeaveAllocation:
        allocation = self.repository.get_by_id(db, allocation_uuid)
        verify_record_ownership(allocation, current_user, detail="Leave allocation not found")
        return allocation

    def create_allocation(
        self, db: Session, data: LeaveAllocationCreate
    ) -> LeaveAllocation:
        employee = self.employee_repository.get_by_id(db, data.employee_id)
        if not employee:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Employee with id '{data.employee_id}' not found",
            )

        leave_type = self.leave_type_repository.get_by_id(db, data.leave_type_id)
        if not leave_type:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Leave type with id '{data.leave_type_id}' not found",
            )

        existing = self.repository.get_by_emp_type_year(
            db,
            employee_id=data.employee_id,
            leave_type_id=data.leave_type_id,
            year=data.year,
            business_id=data.business_id,
        )
        if not existing:
            existing = (
                db.query(LeaveAllocation)
                .filter(
                    LeaveAllocation.employee_id == data.employee_id,
                    LeaveAllocation.leave_type_id == data.leave_type_id,
                    LeaveAllocation.year == data.year,
                )
                .first()
            )

        if existing:
            existing.allocated_days = data.allocated_days
            existing.carried_forward = data.carried_forward
            if data.used_days is not None:
                existing.used_days = data.used_days
            db.commit()
            db.refresh(existing)
            return existing

        return self.repository.create(db, data)

    def update_allocation(
        self, db: Session, allocation_uuid: uuid.UUID, data: LeaveAllocationUpdate, current_user=None
    ) -> LeaveAllocation:
        allocation = self.get_allocation(db, allocation_uuid, current_user=current_user)
        return self.repository.update(db, allocation, data)

    def delete_allocation(self, db: Session, allocation_uuid: uuid.UUID, current_user=None) -> None:
        allocation = self.get_allocation(db, allocation_uuid, current_user=current_user)
        self.repository.delete(db, allocation)

    def generate_export_excel(self, db: Session, business_id: int | None = None) -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Leave Allocations"

        headers = [
            "Allocation ID",
            "Employee ID",
            "Employee Name",
            "Leave Type",
            "Year",
            "Allocated Days",
            "Used Days",
            "Carried Forward",
            "Remaining Balance",
            "Business Profile ID",
            "Business Profile Name",
        ]
        ws.append(headers)

        allocations = self.repository.get_all(db, skip=0, limit=2000, business_id=business_id)
        employees = {e.id: e for e in self.employee_repository.get_all(db, limit=2000)}
        leave_types = {l.id: l.name for l in self.leave_type_repository.get_all(db, limit=1000)}
        biz_repo = BusinessRepository()
        businesses = {b.id: b.name_en for b in biz_repo.get_all(db, skip=0, limit=1000)}

        for a in allocations:
            emp = employees.get(a.employee_id)
            emp_code = emp.employee_id if emp else ""
            emp_name = emp.full_name if emp else ""
            allocated = float(a.allocated_days or 0.0)
            carried = float(a.carried_forward or 0.0)
            used = float(a.used_days or 0.0)
            remaining = allocated + carried - used

            ws.append([
                str(a.id),
                emp_code,
                emp_name,
                leave_types.get(a.leave_type_id, ""),
                a.year,
                allocated,
                used,
                carried,
                remaining,
                str(a.business_id) if a.business_id else "",
                businesses.get(a.business_id, "") if a.business_id else "Global",
            ])

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        output = BytesIO()
        wb.save(output)
        return output.getvalue()


class LeaveApplicationService:
    def __init__(
        self,
        repository: LeaveApplicationRepository | None = None,
        allocation_repository: LeaveAllocationRepository | None = None,
        leave_type_repository: LeaveTypeRepository | None = None,
        employee_repository: EmployeeRepository | None = None,
    ):
        self.repository = repository or LeaveApplicationRepository()
        self.allocation_repository = (
            allocation_repository or LeaveAllocationRepository()
        )
        self.leave_type_repository = leave_type_repository or LeaveTypeRepository()
        self.employee_repository = employee_repository or EmployeeRepository()

    def get_applications(
        self,
        db: Session,
        skip: int = 0,
        limit: int = 100,
        business_id: int | None = None,
        employee_id: uuid.UUID | None = None,
        status_filter: str | LeaveStatusEnum | None = None,
    ) -> list[LeaveApplication]:
        return self.repository.get_all(
            db,
            skip=skip,
            limit=limit,
            business_id=business_id,
            employee_id=employee_id,
            status=status_filter,
        )

    def get_application(
        self, db: Session, application_uuid: uuid.UUID, current_user=None
    ) -> LeaveApplication:
        application = self.repository.get_by_id(db, application_uuid)
        verify_record_ownership(application, current_user, detail="Leave application not found")
        return application

    def _check_overlapping_applications(
        self,
        db: Session,
        employee_id: uuid.UUID,
        start_date,
        end_date,
        exclude_application_id: uuid.UUID | None = None,
    ) -> None:
        existing_apps = self.repository.get_all(
            db, employee_id=employee_id, limit=500
        )
        for app_rec in existing_apps:
            if exclude_application_id and app_rec.id == exclude_application_id:
                continue
            if app_rec.status in (LeaveStatusEnum.PENDING, LeaveStatusEnum.APPROVED):
                if (app_rec.start_date <= end_date) and (app_rec.end_date >= start_date):
                    raise HTTPException(
                        status_code=status.HTTP_400_BAD_REQUEST,
                        detail="Overlapping leave application exists for this date range",
                    )

    def create_application(
        self, db: Session, data: LeaveApplicationCreate
    ) -> LeaveApplication:
        if data.start_date > data.end_date:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Start date cannot be after end date",
            )

        employee = self.employee_repository.get_by_id(db, data.employee_id)
        if not employee:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Employee with id '{data.employee_id}' not found",
            )

        leave_type = self.leave_type_repository.get_by_id(db, data.leave_type_id)
        if not leave_type:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Leave type with id '{data.leave_type_id}' not found",
            )

        # Check overlapping applications
        self._check_overlapping_applications(
            db,
            employee_id=data.employee_id,
            start_date=data.start_date,
            end_date=data.end_date,
        )

        # Compute total_days if 0 or not set
        if data.total_days <= 0:
            data.total_days = (data.end_date - data.start_date).days + 1

        # Check document requirement
        if leave_type.requires_document and not data.document_url:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Supporting document URL is required for this leave type",
            )

        # Check gender applicability
        if leave_type.applicable_gender != GenderApplicabilityEnum.ALL:
            emp_gender = getattr(employee, "gender", None)
            if emp_gender:
                emp_gender_str = (
                    str(emp_gender).value
                    if hasattr(emp_gender, "value")
                    else str(emp_gender)
                )
                if emp_gender_str.lower() != leave_type.applicable_gender.value.lower():
                    raise HTTPException(
                        status_code=status.HTTP_400_BAD_REQUEST,
                        detail=(
                            f"This leave type is only applicable to "
                            f"{leave_type.applicable_gender.value} employees"
                        ),
                    )

        # Check initial balance if leave allocation exists
        if leave_type.max_days_per_year != 0:
            year = data.start_date.year
            allocation = self.allocation_repository.get_by_emp_type_year(
                db,
                employee_id=data.employee_id,
                leave_type_id=data.leave_type_id,
                year=year,
                business_id=data.business_id,
            )
            if allocation:
                remaining = (
                    float(allocation.allocated_days)
                    + float(allocation.carried_forward)
                    - float(allocation.used_days)
                )
                if remaining < data.total_days:
                    raise HTTPException(
                        status_code=status.HTTP_400_BAD_REQUEST,
                        detail=(
                            f"Insufficient leave balance: {remaining} day(s) "
                            f"remaining, {data.total_days} day(s) requested"
                        ),
                    )

        return self.repository.create(db, data)

    def update_application(
        self, db: Session, application_uuid: uuid.UUID, data: LeaveApplicationUpdate, current_user=None
    ) -> LeaveApplication:
        application = self.get_application(db, application_uuid, current_user=current_user)

        if application.status != LeaveStatusEnum.PENDING:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Cannot edit a leave application with status '{application.status}'",
            )

        start_date = data.start_date or application.start_date
        end_date = data.end_date or application.end_date
        if start_date > end_date:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Start date cannot be after end date",
            )

        self._check_overlapping_applications(
            db,
            employee_id=application.employee_id,
            start_date=start_date,
            end_date=end_date,
            exclude_application_id=application_uuid,
        )

        if data.total_days is None or data.total_days <= 0:
            data.total_days = (end_date - start_date).days + 1

        return self.repository.update(db, application, data)

    def review_application(
        self,
        db: Session,
        application_uuid: uuid.UUID,
        data: LeaveApplicationReview,
        current_user=None,
    ) -> LeaveApplication:
        application = self.get_application(db, application_uuid, current_user=current_user)

        if application.status != LeaveStatusEnum.PENDING:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Cannot review leave application with status '{application.status}'",
            )

        reviewer = self.employee_repository.get_by_id(db, data.reviewed_by_id)
        if not reviewer:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Reviewer with id '{data.reviewed_by_id}' not found",
            )

        target_status = (
            data.status
            if isinstance(data.status, LeaveStatusEnum)
            else LeaveStatusEnum(data.status)
        )

        if target_status == LeaveStatusEnum.APPROVED:
            leave_type = self.leave_type_repository.get_by_id(
                db, application.leave_type_id
            )
            if not leave_type:
                raise HTTPException(
                    status_code=status.HTTP_404_NOT_FOUND,
                    detail="Leave type not found for this application",
                )

            # max_days_per_year == 0 means unlimited -> skip balance enforcement
            if leave_type.max_days_per_year != 0:
                year = application.start_date.year
                allocation = self.allocation_repository.get_by_emp_type_year_locked(
                    db,
                    employee_id=application.employee_id,
                    leave_type_id=application.leave_type_id,
                    year=year,
                    business_id=application.business_id,
                )

                if not allocation:
                    raise HTTPException(
                        status_code=status.HTTP_400_BAD_REQUEST,
                        detail=(
                            "No leave allocation configured for this employee, "
                            "leave type, and year. Cannot approve."
                        ),
                    )

                remaining = (
                    float(allocation.allocated_days)
                    + float(allocation.carried_forward)
                    - float(allocation.used_days)
                    - float(allocation.encashed_days)
                )
                if remaining < application.total_days:
                    raise HTTPException(
                        status_code=status.HTTP_400_BAD_REQUEST,
                        detail=(
                            f"Insufficient leave balance: {remaining} day(s) "
                            f"remaining, {application.total_days} day(s) requested"
                        ),
                    )

                allocation.used_days = (
                    float(allocation.used_days) + application.total_days
                )
                db.add(allocation)

            # Bidirectional Attendance Synchronization
            if getattr(data, "sync_attendance", True):
                cur_d = application.start_date
                while cur_d <= application.end_date:
                    att = (
                        db.query(Attendance)
                        .filter(
                            Attendance.business_id == application.business_id,
                            Attendance.employee_id == application.employee_id,
                            Attendance.date == cur_d,
                        )
                        .first()
                    )
                    if att:
                        att.status = AttendanceStatusEnum.ON_LEAVE
                        att.note = f"Approved Leave ({leave_type.code}): {application.reason or ''}".strip()[:250]
                        db.add(att)
                    else:
                        new_att = Attendance(
                            business_id=application.business_id,
                            employee_id=application.employee_id,
                            date=cur_d,
                            status=AttendanceStatusEnum.ON_LEAVE,
                            source=AttendanceSourceEnum.SYSTEM,
                            work_hours=0.0,
                            overtime_hours=0.0,
                            note=f"Approved Leave ({leave_type.code}): {application.reason or ''}".strip()[:250],
                        )
                        db.add(new_att)
                    cur_d += timedelta(days=1)

        # NOTE: allocation and attendance changes above are staged, not committed.
        # update_status() commits them together with the application status
        # change in one transaction. Do not insert a db.commit() between the two.
        return self.repository.update_status(
            db,
            application=application,
            status=target_status,
            reviewed_by_id=data.reviewed_by_id,
            review_note=data.review_note,
        )

    def rollback_application_approval(
        self,
        db: Session,
        application_uuid: uuid.UUID,
        data: LeaveApplicationRollback,
        current_user=None,
    ) -> LeaveApplication:
        """Rollback an APPROVED application back to PENDING status,
        restores the balance, reverts attendance, and logs audit details."""
        application = self.get_application(db, application_uuid, current_user=current_user)

        if application.status != LeaveStatusEnum.APPROVED:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Cannot rollback leave application with status '{application.status}'. Only APPROVED requests can be rolled back.",
            )

        # Check if any date of the leave application falls into a LOCKED or PAID PayrollPeriod
        if not data.force_override:
            locked_period = (
                db.query(PayrollPeriod)
                .filter(
                    PayrollPeriod.business_id == application.business_id,
                    PayrollPeriod.status.in_([PayrollPeriodStatusEnum.LOCKED, PayrollPeriodStatusEnum.PAID]),
                    PayrollPeriod.start_date <= application.end_date,
                    PayrollPeriod.end_date >= application.start_date,
                )
                .first()
            )
            if locked_period:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=(
                        f"Cannot rollback leave approval: This leave period overlaps with a "
                        f"locked or paid payroll period '{locked_period.name}'. "
                        f"Retroactive modification is restricted unless explicitly overridden."
                    ),
                )

        # Restore quota in LeaveAllocation
        leave_type = self.leave_type_repository.get_by_id(db, application.leave_type_id)
        if leave_type and leave_type.max_days_per_year != 0:
            year = application.start_date.year
            allocation = self.allocation_repository.get_by_emp_type_year_locked(
                db,
                employee_id=application.employee_id,
                leave_type_id=application.leave_type_id,
                year=year,
                business_id=application.business_id,
            )
            if allocation:
                restored = float(allocation.used_days) - application.total_days
                allocation.used_days = max(restored, 0)
                db.add(allocation)

        # Rollback synchronized attendance records
        cur_d = application.start_date
        while cur_d <= application.end_date:
            att = (
                db.query(Attendance)
                .filter(
                    Attendance.business_id == application.business_id,
                    Attendance.employee_id == application.employee_id,
                    Attendance.date == cur_d,
                    Attendance.status == AttendanceStatusEnum.ON_LEAVE,
                )
                .first()
            )
            if att:
                if att.source == AttendanceSourceEnum.SYSTEM and (att.work_hours or 0.0) == 0.0:
                    db.delete(att)
                else:
                    att.status = AttendanceStatusEnum.ABSENT
                    att.note = "Leave approval rolled back"
                    db.add(att)
            cur_d += timedelta(days=1)

        return self.repository.rollback_status(
            db,
            application=application,
            rolled_back_by_id=data.actor_id,
            reason=data.reason,
        )

    def cancel_application(
        self,
        db: Session,
        application_uuid: uuid.UUID,
        actor_id: uuid.UUID,
        current_user=None,
    ) -> LeaveApplication:
        """Cancel a PENDING or APPROVED application. If it was APPROVED,
        restores the balance that was deducted at approval time."""
        application = self.get_application(db, application_uuid, current_user=current_user)

        if application.status not in (
            LeaveStatusEnum.PENDING,
            LeaveStatusEnum.APPROVED,
        ):
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Cannot cancel a leave application with status '{application.status}'",
            )

        was_approved = application.status == LeaveStatusEnum.APPROVED

        if was_approved:
            leave_type = self.leave_type_repository.get_by_id(
                db, application.leave_type_id
            )
            if leave_type and leave_type.max_days_per_year != 0:
                year = application.start_date.year
                allocation = self.allocation_repository.get_by_emp_type_year_locked(
                    db,
                    employee_id=application.employee_id,
                    leave_type_id=application.leave_type_id,
                    year=year,
                    business_id=application.business_id,
                )
                if allocation:
                    restored = float(allocation.used_days) - application.total_days
                    allocation.used_days = max(restored, 0)
                    db.add(allocation)
                # If the allocation row no longer exists, there's nothing to
                # restore to -- proceed with cancellation regardless.

        return self.repository.update_status(
            db,
            application=application,
            status=LeaveStatusEnum.CANCELLED,
            reviewed_by_id=actor_id,
            review_note="Cancelled" if was_approved else None,
        )

    def delete_application(self, db: Session, application_uuid: uuid.UUID, current_user=None) -> None:
        application = self.get_application(db, application_uuid, current_user=current_user)
        self.repository.delete(db, application)

    def generate_export_excel(self, db: Session, business_id: int | None = None) -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Leave Requests"

        headers = [
            "Application ID",
            "Employee ID",
            "Employee Name",
            "Leave Type",
            "Start Date",
            "End Date",
            "Total Days",
            "Reason",
            "Status",
            "Review Note",
            "Business Profile ID",
            "Business Profile Name",
        ]
        ws.append(headers)

        apps = self.repository.get_all(db, skip=0, limit=2000, business_id=business_id)
        employees = {e.id: e for e in self.employee_repository.get_all(db, limit=2000)}
        leave_types = {l.id: l.name for l in self.leave_type_repository.get_all(db, limit=1000)}
        biz_repo = BusinessRepository()
        businesses = {b.id: b.name_en for b in biz_repo.get_all(db, skip=0, limit=1000)}

        for a in apps:
            emp = employees.get(a.employee_id)
            emp_code = emp.employee_id if emp else ""
            emp_name = emp.full_name if emp else ""

            ws.append([
                str(a.id),
                emp_code,
                emp_name,
                leave_types.get(a.leave_type_id, ""),
                a.start_date.strftime("%Y-%m-%d") if a.start_date else "",
                a.end_date.strftime("%Y-%m-%d") if a.end_date else "",
                a.total_days or 0,
                a.reason or "",
                a.status.value if a.status else "",
                a.review_note or "",
                str(a.business_id) if a.business_id else "",
                businesses.get(a.business_id, "") if a.business_id else "Global",
            ])

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        output = BytesIO()
        wb.save(output)
        return output.getvalue()

    def get_leave_summary_report(
        self, db: Session, business_id: int, period_start, period_end
    ) -> dict:
        # Added for MCP read-only report aggregation
        from datetime import datetime
        if isinstance(period_start, str):
            period_start_dt = datetime.strptime(period_start, "%Y-%m-%d").date()
        else:
            period_start_dt = period_start
        if isinstance(period_end, str):
            period_end_dt = datetime.strptime(period_end, "%Y-%m-%d").date()
        else:
            period_end_dt = period_end

        apps = self.repository.get_all(db, business_id=business_id, limit=5000)
        filtered_apps = [
            a for a in apps
            if a.start_date <= period_end_dt and a.end_date >= period_start_dt
        ]

        summary_by_type: dict[str, dict[str, int]] = {}
        for app in filtered_apps:
            lt_name = app.leave_type.name if app.leave_type else "Unknown"
            if lt_name not in summary_by_type:
                summary_by_type[lt_name] = {
                    "pending": 0, "approved": 0, "rejected": 0, "cancelled": 0, "total": 0
                }
            st_str = str(app.status.value if hasattr(app.status, "value") else app.status).lower()
            if st_str in summary_by_type[lt_name]:
                summary_by_type[lt_name][st_str] += 1
            summary_by_type[lt_name]["total"] += 1

        return {
            "business_id": business_id,
            "period_start": str(period_start_dt),
            "period_end": str(period_end_dt),
            "total_applications": len(filtered_apps),
            "by_leave_type": summary_by_type,
        }


class LeaveEncashmentService:
    def __init__(
        self,
        repository: LeaveEncashmentRepository | None = None,
        allocation_repository: LeaveAllocationRepository | None = None,
        leave_type_repository: LeaveTypeRepository | None = None,
        employee_repository: EmployeeRepository | None = None,
        salary_repository: EmployeeSalaryRepository | None = None,
    ):
        self.repository = repository or LeaveEncashmentRepository()
        self.allocation_repository = (
            allocation_repository or LeaveAllocationRepository()
        )
        self.leave_type_repository = leave_type_repository or LeaveTypeRepository()
        self.employee_repository = employee_repository or EmployeeRepository()
        self.salary_repository = salary_repository or EmployeeSalaryRepository()

    def get_encashments(
        self,
        db: Session,
        skip: int = 0,
        limit: int = 100,
        business_id: int | None = None,
        employee_id: uuid.UUID | None = None,
        year: int | None = None,
        status: str | LeaveEncashmentStatusEnum | None = None,
    ) -> list[LeaveEncashment]:
        return self.repository.get_all(
            db,
            skip=skip,
            limit=limit,
            business_id=business_id,
            employee_id=employee_id,
            year=year,
            status=status,
        )

    def get_encashment(
        self, db: Session, encashment_uuid: uuid.UUID, current_user=None
    ) -> LeaveEncashment:
        encashment = self.repository.get_by_id(db, encashment_uuid)
        verify_record_ownership(encashment, current_user, detail="Leave encashment not found")
        return encashment

    def create_encashment(
        self, db: Session, data: LeaveEncashmentCreate
    ) -> LeaveEncashment:
        employee = self.employee_repository.get_by_id(db, data.employee_id)
        if not employee:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Employee with id '{data.employee_id}' not found",
            )

        leave_type = self.leave_type_repository.get_by_id(db, data.leave_type_id)
        if not leave_type:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Leave type with id '{data.leave_type_id}' not found",
            )

        # Under BLA 2006, only cumulative / annual leave can be encashed (or leave types explicitly marked carry_forward)
        if not leave_type.carry_forward and leave_type.code not in ("AL", "EL"):
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=(
                    f"Leave type '{leave_type.name}' ({leave_type.code}) is not eligible for encashment. "
                    f"Only cumulative / annual leave (BLA 2006 § 117) may be encashed."
                ),
            )

        # Check balance
        allocation = self.allocation_repository.get_by_emp_type_year_locked(
            db,
            employee_id=data.employee_id,
            leave_type_id=data.leave_type_id,
            year=data.year,
            business_id=data.business_id,
        )
        if not allocation:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="No leave allocation configured for this employee, leave type, and year.",
            )

        remaining = allocation.remaining_days
        if remaining < data.encashable_days:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=(
                    f"Insufficient leave balance: {remaining} day(s) remaining, "
                    f"{data.encashable_days} day(s) requested for encashment."
                ),
            )

        # Calculate daily rate from employee salary (basic_salary / 30.0 per BLA 2006 § 117)
        salary_rec = self.salary_repository.get_by_employee_id(
            db, employee_id=data.employee_id, business_id=data.business_id
        )
        if salary_rec and float(salary_rec.basic_salary or 0.0) > 0.0:
            daily_rate = round(float(salary_rec.basic_salary) / 30.0, 2)
        elif salary_rec and float(salary_rec.gross_salary or 0.0) > 0.0:
            daily_rate = round(float(salary_rec.gross_salary) / 30.0, 2)
        else:
            daily_rate = 500.0  # standard fallback minimum rate

        total_amount = round(data.encashable_days * daily_rate, 2)

        return self.repository.create(
            db,
            data=data,
            daily_rate=daily_rate,
            total_amount=total_amount,
        )

    def review_encashment(
        self,
        db: Session,
        encashment_uuid: uuid.UUID,
        data: LeaveEncashmentReview,
        current_user=None,
    ) -> LeaveEncashment:
        encashment = self.get_encashment(db, encashment_uuid, current_user=current_user)

        if encashment.status != LeaveEncashmentStatusEnum.PENDING:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Cannot review leave encashment with status '{encashment.status}'.",
            )

        reviewer = self.employee_repository.get_by_id(db, data.reviewed_by_id)
        if not reviewer:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Reviewer with id '{data.reviewed_by_id}' not found",
            )

        target_status = (
            data.status
            if isinstance(data.status, LeaveEncashmentStatusEnum)
            else LeaveEncashmentStatusEnum(data.status)
        )

        if target_status == LeaveEncashmentStatusEnum.APPROVED:
            allocation = self.allocation_repository.get_by_emp_type_year_locked(
                db,
                employee_id=encashment.employee_id,
                leave_type_id=encashment.leave_type_id,
                year=encashment.year,
                business_id=encashment.business_id,
            )
            if not allocation:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="Leave allocation row not found for encashment deduction.",
                )

            remaining = allocation.remaining_days
            if remaining < float(encashment.encashable_days):
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=(
                        f"Cannot approve: Insufficient balance. Remaining: {remaining} day(s), "
                        f"requested: {encashment.encashable_days} day(s)."
                    ),
                )

            allocation.encashed_days = (
                float(allocation.encashed_days) + float(encashment.encashable_days)
            )
            db.add(allocation)

        return self.repository.update_status(
            db,
            encashment=encashment,
            status=target_status,
            reviewed_by_id=data.reviewed_by_id,
            review_note=data.review_note,
        )

    def disburse_encashment(
        self,
        db: Session,
        encashment_uuid: uuid.UUID,
        data: LeaveEncashmentDisburse,
        current_user=None,
    ) -> LeaveEncashment:
        encashment = self.get_encashment(db, encashment_uuid, current_user=current_user)

        if encashment.status != LeaveEncashmentStatusEnum.APPROVED:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Cannot disburse encashment with status '{encashment.status}'. Must be APPROVED first.",
            )

        encashment.status = LeaveEncashmentStatusEnum.PAID
        encashment.payment_date = data.payment_date
        encashment.payment_reference = data.payment_reference
        encashment.payroll_period_id = data.payroll_period_id
        db.commit()
        db.refresh(encashment)
        return encashment

    def delete_encashment(
        self, db: Session, encashment_uuid: uuid.UUID, current_user=None
    ) -> None:
        encashment = self.get_encashment(db, encashment_uuid, current_user=current_user)

        if encashment.status == LeaveEncashmentStatusEnum.APPROVED:
            # Restore quota before delete
            allocation = self.allocation_repository.get_by_emp_type_year_locked(
                db,
                employee_id=encashment.employee_id,
                leave_type_id=encashment.leave_type_id,
                year=encashment.year,
                business_id=encashment.business_id,
            )
            if allocation:
                restored = float(allocation.encashed_days) - float(encashment.encashable_days)
                allocation.encashed_days = max(restored, 0)
                db.add(allocation)

        self.repository.delete(db, encashment)


class LeaveCarryForwardService:
    def __init__(
        self,
        allocation_repository: LeaveAllocationRepository | None = None,
        leave_type_repository: LeaveTypeRepository | None = None,
        employee_repository: EmployeeRepository | None = None,
        log_repository: LeaveCarryForwardLogRepository | None = None,
    ):
        self.allocation_repository = (
            allocation_repository or LeaveAllocationRepository()
        )
        self.leave_type_repository = leave_type_repository or LeaveTypeRepository()
        self.employee_repository = employee_repository or EmployeeRepository()
        self.log_repository = log_repository or LeaveCarryForwardLogRepository()

    def preview_carry_forward(
        self,
        db: Session,
        data: LeaveCarryForwardPreviewRequest,
        business_id: int,
    ) -> LeaveCarryForwardPreviewResponse:
        employees = [
            e for e in self.employee_repository.get_all(db, business_id=business_id, limit=2000)
            if e.is_active
        ]
        if data.leave_type_id:
            target_type = self.leave_type_repository.get_by_id(db, data.leave_type_id)
            leave_types = [target_type] if target_type else []
        else:
            leave_types = [
                lt for lt in self.leave_type_repository.get_all(db, business_id=business_id, limit=200)
                if lt.is_active
            ]

        # Map existing allocations for from_year
        allocations_list = self.allocation_repository.get_all(
            db, business_id=business_id, year=data.from_year, limit=5000
        )
        alloc_map = {(a.employee_id, a.leave_type_id): a for a in allocations_list}

        preview_items: list[LeaveCarryForwardPreviewItem] = []
        total_carried = 0.0
        total_lapsed = 0.0

        for emp in employees:
            for lt in leave_types:
                alloc = alloc_map.get((emp.id, lt.id))
                balance = alloc.remaining_days if alloc else 0.0

                if lt.carry_forward:
                    # Apply statutory/custom cap
                    statutory_cap = (
                        data.max_cap
                        if data.max_cap is not None
                        else (float(lt.max_days_per_year * 2) if lt.max_days_per_year > 0 else 60.0)
                    )
                    proposed_carry = min(max(balance, 0.0), statutory_cap)
                    proposed_lapsed = max(balance - proposed_carry, 0.0)
                else:
                    proposed_carry = 0.0
                    proposed_lapsed = max(balance, 0.0)

                new_annual = float(lt.max_days_per_year or 0.0)
                projected_starting = proposed_carry + new_annual

                total_carried += proposed_carry
                total_lapsed += proposed_lapsed

                preview_items.append(
                    LeaveCarryForwardPreviewItem(
                        employee_id=emp.id,
                        employee_code=emp.employee_id or "",
                        employee_name=emp.full_name,
                        department_name=getattr(emp, "department_name", None) or (emp.department.name if getattr(emp, "department", None) else "General"),
                        leave_type_id=lt.id,
                        leave_type_name=lt.name,
                        leave_type_code=lt.code,
                        carry_forward_allowed=lt.carry_forward,
                        from_year_balance=round(balance, 2),
                        proposed_carry_forward=round(proposed_carry, 2),
                        proposed_lapsed=round(proposed_lapsed, 2),
                        new_annual_allocation=round(new_annual, 2),
                        projected_starting_balance=round(projected_starting, 2),
                    )
                )

        return LeaveCarryForwardPreviewResponse(
            from_year=data.from_year,
            to_year=data.to_year,
            total_employees=len(employees),
            total_carried_days=round(total_carried, 2),
            total_lapsed_days=round(total_lapsed, 2),
            items=preview_items,
        )

    def execute_carry_forward(
        self,
        db: Session,
        data: LeaveCarryForwardExecuteRequest,
        business_id: int,
        current_user=None,
    ) -> LeaveCarryForwardExecuteResponse:
        preview_req = LeaveCarryForwardPreviewRequest(
            from_year=data.from_year,
            to_year=data.to_year,
            leave_type_id=data.leave_type_id,
            max_cap=data.max_cap,
            business_id=business_id,
        )
        sim = self.preview_carry_forward(db, preview_req, business_id=business_id)

        count_mutated = 0
        for item in sim.items:
            existing_target = self.allocation_repository.get_by_emp_type_year_locked(
                db,
                employee_id=item.employee_id,
                leave_type_id=item.leave_type_id,
                year=data.to_year,
                business_id=business_id,
            )
            if existing_target:
                existing_target.carried_forward = item.proposed_carry_forward
                existing_target.allocated_days = item.new_annual_allocation
                db.add(existing_target)
            else:
                new_alloc = LeaveAllocation(
                    business_id=business_id,
                    employee_id=item.employee_id,
                    leave_type_id=item.leave_type_id,
                    year=data.to_year,
                    allocated_days=item.new_annual_allocation,
                    carried_forward=item.proposed_carry_forward,
                    used_days=0.0,
                    encashed_days=0.0,
                )
                db.add(new_alloc)
            count_mutated += 1

        # Audit logging
        self.log_repository.create(
            db,
            business_id=business_id,
            from_year=data.from_year,
            to_year=data.to_year,
            leave_type_id=data.leave_type_id,
            total_employees=sim.total_employees,
            total_days_carried=sim.total_carried_days,
            total_days_lapsed=sim.total_lapsed_days,
            executed_by_id=data.executed_by_id,
            notes=data.notes,
        )

        db.commit()

        return LeaveCarryForwardExecuteResponse(
            from_year=data.from_year,
            to_year=data.to_year,
            total_employees=sim.total_employees,
            total_carried_days=sim.total_carried_days,
            total_lapsed_days=sim.total_lapsed_days,
            allocations_created_or_updated=count_mutated,
            message=(
                f"Successfully completed yearly leave carry forward from {data.from_year} to {data.to_year}. "
                f"{count_mutated} quota allocation(s) updated ({sim.total_carried_days} days carried forward, "
                f"{sim.total_lapsed_days} days lapsed)."
            ),
        )

    def get_logs(
        self,
        db: Session,
        skip: int = 0,
        limit: int = 100,
        business_id: int | None = None,
    ) -> list[LeaveCarryForwardLog]:
        return self.log_repository.get_all(
            db, skip=skip, limit=limit, business_id=business_id
        )